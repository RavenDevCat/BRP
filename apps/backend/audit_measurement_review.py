"""Remeasure saved Audit candidates through the existing gates, without solving."""
from __future__ import annotations

from copy import copy, deepcopy
from dataclasses import asdict
from io import BytesIO
import math
from types import SimpleNamespace
from typing import Any, Callable

try:
    from . import planner_core as core
    from . import measurement_reviews as reviews
    from .full_measurement_review import FullReviewProvider, MAX_FULL_ROUTES, _verified
except ImportError:
    import planner_core as core
    import measurement_reviews as reviews
    from full_measurement_review import FullReviewProvider, MAX_FULL_ROUTES, _verified


MODE = "full_audit"
SCENARIOS = ("current_plan", "time_constrained", "exception_preserving")
ALIASES = {"current_plan": "current_plan_scenario", "time_constrained": "time_constrained_optimization",
           "exception_preserving": "exception_preserving_optimization"}
REQUIRED_CONFIG = ("service_direction", "stop_service_minutes", "time_window_start", "time_window_end",
                   "to_school_arrival_time", "from_school_departure_time", "max_route_duration_minutes",
                   "time_impact_limit_minutes", "minimum_vehicle_reduction", "comfort_load_factor", "route_stop_limit")
REVALIDATED_KEYS = ("traffic_gate", "traffic_feasible", "arrival_reverse_check", "final_time_impact_gate",
                    "feasibility_report", "exception_feasible", "scenario_status", "vehicle_saving_target",
                    "time_impact", "decision_metrics", "vehicle_ladder_search", "constraint_search_outcome")


def build_audit_review_request(source: dict[str, Any], *, requested_by: str, request_key: str,
                              provider_call_limit: int) -> dict[str, Any]:
    if (source.get("status") not in {"succeeded", "failed"} or not source.get("result")
            or dict(source.get("metadata") or {}).get("job_kind", "route_audit") != "route_audit"):
        raise ValueError("Full Audit correction requires a finished saved Audit result.")
    if not source.get("job_id") or not requested_by.strip() or not 1 <= len(request_key.strip()) <= 80:
        raise ValueError("Source, administrator and request key are required.")
    if (isinstance(provider_call_limit, bool) or not isinstance(provider_call_limit, int)
            or not 1 <= provider_call_limit <= reviews.MAX_REVIEW_CALLS):
        raise ValueError("Provider call limit must be an integer from 1 to 500.")
    saved = dict(source["result"])
    structured = deepcopy(dict(saved.get("structured_results") or {}))
    if any(key not in structured or not isinstance(structured[key], dict) for key in SCENARIOS):
        raise ValueError("The saved Audit scenario set is incomplete; a full correction cannot omit a plan.")
    original_config = dict(saved.get("planner_config") or structured.get("planner_config") or source.get("config") or {})
    if any(key not in original_config or (original_config[key] is None and key != "route_stop_limit") for key in REQUIRED_CONFIG):
        raise ValueError("Original Audit acceptance parameters are missing; do not infer current defaults.")
    config = core.build_planner_config(original_config)
    for key in ("stop_service_minutes", "max_route_duration_minutes", "time_impact_limit_minutes", "minimum_vehicle_reduction", "comfort_load_factor"):
        if reviews._number(original_config[key]) != getattr(config, key):
            raise ValueError("Original Audit parameters cannot be silently rounded or replaced.")
    current = dict(structured.get("current_plan") or {})
    if not current.get("routes") or not current.get("points"):
        raise ValueError("The saved Current Plan is required as the time-impact baseline.")
    count = 0
    for key in SCENARIOS:
        scenario = structured.get(key)
        if scenario is None:
            continue
        points, routes = list(scenario.get("points") or []), list(scenario.get("routes") or [])
        if not routes:
            continue
        if (not points or points[0].get("is_depot") is not True or any(point.get("is_depot") for point in points[1:]) or any(
                str(point.get("country") or "").lower() != "china" for point in points)):
            raise ValueError("Full Audit correction requires saved China points with school node zero.")
        for point in points:
            riders = reviews._number(point.get("passenger_count"))
            if riders is None or riders != int(riders):
                raise ValueError("Original Audit student counts are missing or invalid.")
        if scenario.get("bus_count", len(routes)) != len(routes):
            raise ValueError("Saved fleet count disagrees with its route list.")
        ids = [core._route_display_id(route, index) for index, route in enumerate(routes, start=1)]
        if not all(ids) or len(set(ids)) != len(ids):
            raise ValueError("Saved route identities are missing or ambiguous.")
        if key == "exception_preserving":
            declared = dict(scenario.get("exception_preserving") or {}).get("frozen_route_ids")
            actual = {route_id for route, route_id in zip(routes, ids) if route.get("exception_role") == "frozen_current"}
            if declared is not None and {str(value) for value in declared} != actual:
                raise ValueError("Saved Protected frozen-route identities disagree.")
        for route, route_id in zip(routes, ids):
            nodes = list(route.get("nodes") or [])
            explicit_id = route.get("route_id") or route.get("id")
            if ((explicit_id and str(explicit_id) != route_id) or len(nodes) < 2 or any(
                    isinstance(node, bool) or not isinstance(node, int) or not 0 <= node < len(points) for node in nodes)
                    or nodes.count(0) != 1 or nodes[-1 if config.service_direction == "To School" else 0] != 0):
                raise ValueError("Saved Audit stop order or school placement is incomplete.")
            if len(route.get("leg_details") or []) != len(nodes) - 1:
                raise ValueError("Original per-leg references are missing; time-impact comparison would be ambiguous.")
            dwell = reviews._number(route.get("stop_service_time_s"))
            # Audit charges dwell on arrival at non-school nodes, not at the AM origin.
            dwell_stops = sum(node != 0 for node in nodes[1:])
            if dwell is None or not math.isclose(dwell, dwell_stops * config.stop_service_minutes * 60, abs_tol=.01):
                raise ValueError("Saved route dwell disagrees with the original parameters.")
            if reviews._number(route.get("load")) is None or not reviews._number(route.get("bus_capacity")):
                raise ValueError("Saved route load or physical capacity is missing.")
        count += len(routes)
    if not 1 <= count <= MAX_FULL_ROUTES:
        raise ValueError("Full Audit correction requires 1-200 saved routes across its scenarios.")
    routes = reviews.source_route_scopes({**source, "result": {**saved, "structured_results": structured}, "config": asdict(config)})
    if len(routes) != count:
        raise ValueError("Every saved Audit route must be included in the AMap correction.")
    inputs = {"structured_results": structured, "config": asdict(config), "input_records": deepcopy(current["points"])}
    return {"review_version": reviews.REVIEW_VERSION, "evidence_version": reviews.EVIDENCE_VERSION, "mode": MODE,
        "source_job_id": str(source["job_id"]), "owner_email": str(source.get("owner_email") or "").strip().lower(),
        "requested_by": requested_by.strip().lower(), "request_key": request_key,
        "source_finished_at": source.get("finished_at"), "source_result_digest": reviews._digest(saved),
        "input_digest": reviews._digest(routes), "routes": routes, "provider_call_limit": provider_call_limit,
        "full_input": inputs, "full_input_digest": reviews._digest(inputs),
        "scope_summary": {"route_count": count, "scenario_count": sum(bool(structured.get(key)) for key in SCENARIOS)}}


def _clear_derived(scenario: dict[str, Any]) -> None:
    scenario["source_acceptance"] = {key: deepcopy(scenario.get(key)) for key in REVALIDATED_KEYS if key in scenario}
    for key in (*REVALIDATED_KEYS, "output_html", "infeasible_reason", "unresolved_reason"):
        scenario.pop(key, None)
    exception = dict(scenario.get("exception_preserving") or {})
    scenario["source_exception_preserving"] = deepcopy(exception)
    for key in ("accepted", "candidate_failure_summary", "candidate_remainder_failure_summary", "current_failure_summary", "attempts"):
        exception.pop(key, None)
    if "exception_preserving" in scenario:
        scenario["exception_preserving"] = exception
    constraint = dict(scenario.get("time_constraint") or {})
    for key in ("final_validation_required", "final_validation_status", "final_satisfied"):
        constraint.pop(key, None)
    if "time_constraint" in scenario:
        scenario["time_constraint"] = constraint
    for route in scenario.get("routes") or []:
        for key in ("route_evidence", "final_route_traffic_gate", "am_arrival_gate", "arrival_reverse_check", "time_impact"):
            route.pop(key, None)
    for point in scenario.get("points") or []:
        point.pop("time_impact", None)


def run_audit_review(store: Any, record: dict[str, Any], token: str, *,
                     provider_factory: Callable, checkpoint: Callable) -> dict[str, Any]:
    request, rid = record["request"], record["review_id"]
    if (request.get("review_version") != reviews.REVIEW_VERSION or request.get("evidence_version") != reviews.EVIDENCE_VERSION
            or reviews._digest(request.get("full_input")) != request.get("full_input_digest")
            or reviews._digest(request.get("routes")) != request.get("input_digest")):
        raise ValueError("Full Audit input or measurement contract changed.")
    def check_active():
        if not store.route_measurement_claim_active(rid, token):
            raise reviews.ReviewClaimLost("Audit review is no longer running.")
    check_active()
    used, limit = record.get("api_calls", 0), request.get("provider_call_limit")
    if (isinstance(limit, bool) or not isinstance(limit, int) or not 1 <= limit <= reviews.MAX_REVIEW_CALLS
            or isinstance(used, bool) or not isinstance(used, int) or not 0 <= used <= limit):
        raise ValueError("Invalid persisted provider budget or usage.")
    config = core.build_planner_config(request["full_input"]["config"])
    inner = provider_factory("amap", departure_time=None, api_call_limit=request["provider_call_limit"])
    inner.state = reviews.ReviewCallState({**inner.state, "api_calls": int(record.get("api_calls") or 0)},
        lambda count: store.reserve_route_measurement_calls(rid, token, count))
    provider = FullReviewProvider(inner,
        read_snapshot=lambda key: store.get_review_measurement_snapshot(rid, token, key),
        write_snapshot=lambda key, value: store.save_review_measurement_snapshot(rid, token, key, value), check_active=check_active)
    route_limit_s = core.effective_route_duration_limit_minutes(config) * 60
    planner = SimpleNamespace(AMAP_KEY=getattr(inner.planner, "AMAP_KEY", ""),
        MAX_ROUTE_DURATION_SECONDS=route_limit_s,
        _BRP_FINAL_ROUTE_TRAFFIC_GATE_DURATION_SECONDS=route_limit_s,
        _BRP_RUNTIME_PROFILE={})
    structured = deepcopy(request["full_input"]["structured_results"])
    for key in ("output_paths", "current_plan_comparison", "route_reallocation_analysis", "input_address_review", "runtime_profile", "scheduled_run_traffic_refresh"):
        if key in structured:
            structured.setdefault("source_reference_context", {})[key] = structured.pop(key)
    structured["job_id"] = rid
    output = {"review_version": reviews.REVIEW_VERSION, "mode": MODE, "scope": "full_audit_result",
        "source_job_id": request["source_job_id"], "source_result_digest": request["source_result_digest"],
        "full_input_digest": request["full_input_digest"], "status": "running", "routes": [],
        "provider_api_calls": int(record.get("api_calls") or 0), "audit_result": {},
        "comparison_basis": "fresh_traffic_not_controlled_before_after", "time_window_revalidated": False,
        "time_impact_revalidated": False, "solver_rerun": False, "minimum_fleet_revalidated": False}
    for key in SCENARIOS:
        if key in structured:
            _clear_derived(structured[key])
    for key in SCENARIOS:
        scenario = structured.get(key)
        if not scenario:
            continue
        check_active()
        if not scenario.get("routes"):
            scenario.update(scenario_status="unresolved", unresolved_reason="No saved candidate to remeasure.")
            continue
        core.attach_final_route_traffic_gate(planner, scenario, scenario["points"], config,
            request["full_input"]["input_records"], key, measurement_provider=provider, check_canceled=check_active)
        if scenario["traffic_gate"].get("status") in {"disabled", "not_applicable"}:
            scenario["traffic_gate"].update(status="unavailable", reason="correction_requires_live_measurements",
                unavailable_route_count=len(scenario["routes"]))
            scenario["traffic_feasible"] = False
        for route in scenario["routes"]:
            route.setdefault("final_route_traffic_gate", {"status": "unavailable", "passes": None})
            if not route.get("route_evidence"):
                route["route_evidence"] = {"evidence_version": reviews.EVIDENCE_VERSION, "provider": "amap",
                    "status": "needs_review", "complete": False, "geometry": [], "legs": [],
                    "duration_s": None, "distance_m": None, "issues": [{"code": "measurement_unavailable"}]}
        output["provider_api_calls"] = int(provider.state["api_calls"])
        output["audit_result"] = _native_result(structured, config)
        checkpoint(deepcopy(output))
    current = structured["current_plan"]
    validate_impact = core._build_final_time_impact_validator(current, current["points"], config, config.time_impact_limit_minutes)
    current_count = len(current["routes"])
    target = max(0, current_count - config.minimum_vehicle_reduction)
    for key in SCENARIOS:
        scenario = structured.get(key)
        if not scenario or not scenario.get("routes"):
            continue
        frozen = {core._route_display_id(route, index) for index, route in enumerate(scenario["routes"], start=1)
                  if route.get("exception_role") == "frozen_current"} if key == "exception_preserving" else set()
        if key != "current_plan":
            validate_impact(scenario, scenario["points"])
        full_gate = deepcopy(scenario["traffic_gate"])
        gate = core._scope_protected_traffic_gate(scenario, frozen) if key == "exception_preserving" else full_gate
        scenario["feasibility_report"] = core.build_route_feasibility_report(scenario, gate, config,
            max_vehicle_count=target if key != "current_plan" else 0, ignored_route_ids=frozen)
        if key != "current_plan":
            core._apply_vehicle_saving_target(scenario, current_count, config.minimum_vehicle_reduction, frozen_route_count=len(frozen))
        accepted = core._scenario_feasibility_passed(scenario)
        if key == "exception_preserving":
            summary = core._scenario_exception_summary(scenario, config=config)
            remainder = core._scenario_exception_summary(scenario, include_frozen_current=False, config=config)
            accepted &= core._exception_candidate_accepted(remainder, len(scenario["routes"]), current_count, target)
            scenario["exception_preserving"] = {**dict(scenario.get("exception_preserving") or {}), "accepted": accepted,
                "frozen_route_count": len(frozen), "frozen_route_ids": sorted(frozen),
                "current_failure_summary": core._scenario_exception_summary(current, config=config),
                "candidate_failure_summary": summary, "candidate_remainder_failure_summary": remainder}
            scenario["exception_feasible"] = accepted
        scenario["scenario_status"] = "passed" if accepted else "unresolved" if full_gate["status"] in {"disabled", "unavailable", "not_applicable"} else "rejected"
        scenario["constraint_search_outcome"] = {"status": scenario["scenario_status"], "reason": "saved_candidate_remeasured_without_solver",
            "search_complete": False, "minimum_fleet_revalidated": False, "attempted_vehicle_caps": []}
        measured = [dict(route.get("final_route_traffic_gate") or {}) for route in scenario["routes"]]
        complete = all(g.get("status") in {"passed", "failed"} for g in measured)
        scenario["planning_reference_metrics"] = {field: scenario.get(field) for field in ("avg_route_duration_s", "avg_route_distance_m")}
        scenario["avg_route_duration_s"] = sum(g["verified_total_duration_s"] for g in measured) / len(measured) if complete else None
        scenario["avg_route_distance_m"] = sum(g["verified_distance_m"] for g in measured) / len(measured) if complete else None
        sources = sorted({str((route.get("route_evidence") or {}).get("source") or "unavailable") for route in scenario["routes"]})
        scenario["measurement_summary"] = {"complete": complete, "measured_route_count": sum(g.get("status") in {"passed", "failed"} for g in measured),
            "route_count": len(measured), "source": sources[0] if len(sources) == 1 else "mixed_amap_evidence",
            "sources": sources, "financial_basis": "unchanged_planning_reference"}
    output["audit_result"] = _native_result(structured, config)
    for scope in request["routes"]:
        scenario_key = scope["route_key"].split(":", 1)[0]
        route = next(row for index, row in enumerate(structured[scenario_key]["routes"], start=1)
                     if core._route_display_id(row, index) == scope["route_id"])
        evidence, gate = dict(route.get("route_evidence") or {}), dict(route.get("final_route_traffic_gate") or {})
        certified = _verified(evidence, len(scope["points"])) and gate.get("status") in {"passed", "failed"}
        after = {"drive_duration_s": gate.get("verified_drive_duration_s") if certified else None,
            "distance_m": gate.get("verified_distance_m") if certified else None,
            "total_duration_s": gate.get("verified_total_duration_s") if certified else None, "captured_at": evidence.get("called_at")}
        output["routes"].append({"route_key": scope["route_key"], "route_id": scope["route_id"], "vehicle_id": scope["vehicle_id"],
            "status": "verified" if certified else "needs_review", "points": scope["points"], "before": scope["before"],
            "after": after, "route_evidence": evidence, "delta": {key: after[key] - scope["before"][key]
                if after[key] is not None and scope["before"].get(key) is not None else None
                for key in ("drive_duration_s", "distance_m", "total_duration_s")}})
    saved_scenarios = [structured[key] for key in SCENARIOS if key in structured]
    output["time_window_revalidated"] = all(s.get("routes") and s.get("traffic_gate", {}).get("status") in {"passed", "failed"}
        and not s["traffic_gate"].get("unavailable_route_count") for s in saved_scenarios)
    output["time_impact_revalidated"] = all(structured[key].get("final_time_impact_gate", {}).get("status") in {"passed", "failed"}
        for key in SCENARIOS[1:] if key in structured)
    output["status"] = "complete" if output["time_window_revalidated"] and output["time_impact_revalidated"] and all(
        row["status"] == "verified" for row in output["routes"]) else "partial"
    output["scenario_changes"] = {key: {"before": structured[key].get("source_acceptance", {}),
        "after": {field: structured[key].get(field) for field in REVALIDATED_KEYS if field in structured[key]}}
        for key in SCENARIOS if key in structured}
    check_active()
    return output


def _native_result(structured: dict[str, Any], config: core.PlannerConfig) -> dict[str, Any]:
    for key, alias in ALIASES.items():
        if key in structured:
            structured[alias] = deepcopy(structured[key])
    structured["planner_config"] = asdict(config)
    structured["service_direction"] = config.service_direction
    assessment = deepcopy(dict(structured.get("current_plan_assessment") or {}))
    current = structured.get("current_plan", {})
    if current.get("measurement_summary"):
        assessment["avg_route_duration_s"] = current.get("avg_route_duration_s")
        assessment["avg_route_distance_m"] = current.get("avg_route_distance_m")
        routes = {core._route_display_id(route, index): route for index, route in enumerate(current.get("routes") or [], start=1)}
        gates = [route.get("final_route_traffic_gate", {}) for route in routes.values()]
        complete = bool(current["measurement_summary"]["complete"])
        assessment["total_duration_s"] = sum(g["verified_total_duration_s"] for g in gates) if complete else None
        assessment["total_distance_m"] = sum(g["verified_distance_m"] for g in gates) if complete else None
        assessment["overlong_route_count"] = sum(g["verified_total_duration_s"] > core.effective_route_duration_limit_minutes(config) * 60 for g in gates) if complete else None
        if assessment.get("recommendations"):
            structured.setdefault("source_reference_context", {})["current_plan_recommendations"] = assessment.pop("recommendations")
        for row in assessment.get("route_summaries") or []:
            gate = routes.get(row.get("route_id"), {}).get("final_route_traffic_gate", {})
            row["duration_s"], row["distance_m"] = gate.get("verified_total_duration_s"), gate.get("verified_distance_m")
        structured["current_plan_assessment"] = assessment
    return {"structured_results": deepcopy(structured), "planner_config": asdict(config), "current_plan_assessment": assessment,
        "service_direction": config.service_direction, "summary": core.summarize_structured_results(structured, int(structured.get("input_address_count") or 0)),
        **{alias: deepcopy(structured[key]) for key, alias in ALIASES.items() if key in structured}}


def validate_audit_result(request: dict[str, Any], result: dict[str, Any], *, terminal: bool) -> None:
    if result.get("status") == "failed":
        return
    if result.get("scope") != "full_audit_result" or result.get("full_input_digest") != request["full_input_digest"]:
        raise ValueError("Full Audit correction does not match its input snapshot.")
    native = dict(result.get("audit_result") or {})
    if native.get("planner_config") != request["full_input"]["config"]:
        raise ValueError("Full Audit correction must retain original parameters.")
    if terminal and result.get("status") in {"complete", "partial"}:
        original, corrected = request["full_input"]["structured_results"], native.get("structured_results", {})
        for key in SCENARIOS:
            if (key in original) != (key in corrected):
                raise ValueError("Full Audit correction must retain every saved scenario.")
            if key not in original:
                continue
            if corrected.get(ALIASES[key]) != corrected[key] or native.get(ALIASES[key]) != corrected[key]:
                raise ValueError("Corrected Audit scenario aliases must use the same result.")
            def identity(scenario):
                return [{field: route.get(field) for field in ("route_id", "vehicle_id", "nodes", "load", "bus_capacity", "exception_role", "stop_service_time_s")}
                        for route in scenario.get("routes") or []]
            if identity(original[key]) != identity(corrected[key]):
                raise ValueError("Full Audit correction cannot change routes, vehicles, students or frozen roles.")
            old_points, new_points = deepcopy(original[key].get("points", [])), deepcopy(corrected[key].get("points", []))
            for point in [*old_points, *new_points]:
                point.pop("time_impact", None)
            if old_points != new_points:
                raise ValueError("Full Audit correction cannot replace saved pickup points.")


def audit_review_record(record: dict[str, Any]) -> dict[str, Any]:
    result = dict(record.get("result") or {})
    if (record.get("status") not in {"succeeded", "needs_review"} or result.get("scope") != "full_audit_result"
            or not result.get("audit_result")):
        raise ValueError("A finalized full Audit correction is required.")
    native = deepcopy(result["audit_result"])
    return {"job_id": record["review_id"], "config": native["planner_config"], "result": native,
            "metadata": {"job_kind": "route_audit", "job_name": f"Measurement correction of {record['source_job_id']}"}}


def build_audit_review_workbook(record: dict[str, Any], *, export_time_impact: Callable) -> bytes:
    from openpyxl import Workbook, load_workbook
    try:
        from .direct_school_analysis import _write_readable_table
    except ImportError:
        from direct_school_analysis import _write_readable_table
    native = audit_review_record(record)
    result = record["result"]
    names = {"current_plan": "Current Plan", "time_constrained": "Strict Plan", "exception_preserving": "Protected Plan"}
    statuses = {"passed": "Passed", "failed": "Not passed", "unavailable": "Needs review", "unresolved": "Needs review"}
    scopes = {"all_routes": "All routes", "optimized_remainder": "Non-frozen routes"}
    book = Workbook()
    book.active.title = "Acceptance Review"
    rows = []
    for key in SCENARIOS:
        change = result.get("scenario_changes", {}).get(key)
        if change is None:
            continue
        before, after = change["before"], change["after"]
        gate = dict(after.get("traffic_gate") or {})
        impact = dict(after.get("final_time_impact_gate") or {})
        rows.append([names[key], statuses.get(dict(before.get("feasibility_report") or {}).get("status"), "Unavailable"),
            statuses.get(dict(after.get("feasibility_report") or {}).get("status"), "Needs review"),
            statuses.get(gate.get("all_routes_status", gate.get("status")), "Needs review"),
            scopes.get(gate.get("evaluation_scope", "all_routes"), "Needs review"),
            "Baseline" if key == "current_plan" else statuses.get(impact.get("status"), "Needs review"), impact.get("over_limit_rider_count"),
            impact.get("max_adverse_minutes")])
    _write_readable_table(book.active, "Audit measurement correction",
        "Fresh traffic is not a controlled before/after test. Original routes are unchanged. Minimum fleet size has not been revalidated.",
        ["Plan", "Original acceptance", "Corrected acceptance", "All-route window", "Acceptance scope", "Student time impact", "Over-limit students", "Max adverse min"],
        rows, [23, 25, 25, 25, 27, 25, 24, 24])
    def minutes(value):
        return value / 60 if value is not None else None
    def kilometres(value):
        return value / 1000 if value is not None else None
    display_names = {}
    scenarios = native["result"]["structured_results"]
    for key in SCENARIOS:
        view = core.attach_route_display_metadata(deepcopy(scenarios[key]), optimized=key != "current_plan",
            reference_routes=scenarios["current_plan"]["routes"] if key == "time_constrained" else None)
        for index, route in enumerate(view.get("routes") or [], start=1):
            display_names[f"{key}:{core._route_display_id(route, index)}"] = route["display_route_id"]
    comparisons = [[names[row["route_key"].split(":", 1)[0]], display_names[row["route_key"]],
        minutes(row["before"].get("total_duration_s")), minutes(row["after"].get("total_duration_s")),
        minutes(row["delta"].get("total_duration_s")), kilometres(row["before"].get("distance_m")),
        kilometres(row["after"].get("distance_m")), row["status"], row["after"].get("captured_at")]
        for row in result.get("routes") or []]
    _write_readable_table(book.create_sheet("Road Comparison"), "Original and corrected road measurements",
        "Blank values mean unavailable evidence. Protected frozen routes remain in this comparison even when excluded from remainder acceptance.",
        ["Plan", "Route", "Original min", "Corrected min", "Change min", "Original km", "Corrected km", "Evidence", "Measured at"],
        comparisons, [23, 24, 20, 20, 20, 20, 20, 22, 44])
    for row in book["Road Comparison"].iter_rows(min_row=5, min_col=3, max_col=7):
        for cell in row:
            cell.number_format = "0.0"
    for key, prefix in (("time_constrained", "Strict"), ("exception_preserving", "Protected")):
        if key not in native["result"]["structured_results"]:
            continue
        scenarios = native["result"]["structured_results"]
        if all(scenarios.get(name, {}).get("measurement_summary", {}).get("complete") for name in ("current_plan", key)):
            content, error = export_time_impact(native, key)
        else:
            content, error = None, "Corrected route or Current Plan baseline measurements are incomplete."
        if not content:
            _write_readable_table(book.create_sheet(f"{prefix} Availability"), f"{names[key]} time-impact evidence",
                "No complete student comparison is asserted when measurement evidence is missing.", ["Status", "Reason"],
                [["Needs review", error or "Time-impact evidence unavailable."]], [24, 95])
            continue
        source_book = load_workbook(BytesIO(content))
        for sheet in source_book:
            target = book.create_sheet(f"{prefix} {sheet.title}"[:31])
            for source_row in sheet:
                for cell in source_row:
                    dest = target.cell(cell.row, cell.column, cell.value)
                    if cell.has_style:
                        dest.font, dest.fill, dest.border = copy(cell.font), copy(cell.fill), copy(cell.border)
                        dest.alignment, dest.protection, dest.number_format = copy(cell.alignment), copy(cell.protection), cell.number_format
            for key, dimension in sheet.column_dimensions.items():
                target.column_dimensions[key] = copy(dimension)
            for key, dimension in sheet.row_dimensions.items():
                target.row_dimensions[key] = copy(dimension)
            for merged in sheet.merged_cells.ranges:
                target.merge_cells(str(merged))
            target.freeze_panes, target.auto_filter.ref = sheet.freeze_panes, sheet.auto_filter.ref
    output = BytesIO()
    book.save(output)
    return output.getvalue()
