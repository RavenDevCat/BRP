"""Append full Direct-to-School corrections using the existing analysis pipeline."""
from __future__ import annotations

from collections import Counter
from copy import deepcopy
from datetime import datetime, timezone
from io import BytesIO
import math
import re
from typing import Any, Callable

from openpyxl import load_workbook

try:
    from . import direct_school_analysis as analysis
    from . import measurement_reviews as reviews
    from .route_evidence import CACHE_MAX_AGE_SECONDS, EVIDENCE_VERSION
except ImportError:
    import direct_school_analysis as analysis
    import measurement_reviews as reviews
    from route_evidence import CACHE_MAX_AGE_SECONDS, EVIDENCE_VERSION


MAX_FULL_ROUTES = 200
MODE = "full_direct_school"


def build_full_review_request(source: dict[str, Any], *, requested_by: str,
                              request_key: str, provider_call_limit: int) -> dict[str, Any]:
    if (source.get("status") not in {"succeeded", "failed"} or not source.get("result")
            or dict(source.get("metadata") or {}).get("job_kind") != "direct_school_analysis"):
        raise ValueError("Full correction requires a finished Direct-to-School result.")
    if not source.get("job_id") or not requested_by.strip() or not 1 <= len(request_key.strip()) <= 80:
        raise ValueError("Source, requesting administrator and request key are required.")
    if (isinstance(provider_call_limit, bool) or not isinstance(provider_call_limit, int)
            or not 1 <= provider_call_limit <= reviews.MAX_REVIEW_CALLS):
        raise ValueError("Provider call limit must be an integer from 1 to 500.")
    saved = dict(source["result"])
    config = dict(saved.get("parameters") or {})
    if config.get("service_direction") not in {"To School", "From School"}:
        raise ValueError("The original service direction is missing.")
    for key in ("far_duration_minutes", "stop_service_minutes"):
        if reviews._number(config.get(key)) is None:
            raise ValueError(f"The original {key} is missing or invalid; do not infer a default.")
    clock_keys = ["time_window_start", "time_window_end"]
    if config["service_direction"] == "From School":
        clock_keys.append("from_school_departure_time")
    for key in clock_keys:
        if not re.fullmatch(r"(?:[01]?\d|2[0-3]):[0-5]\d", str(config.get(key) or "")):
            raise ValueError(f"The original {key} is missing or invalid; do not infer a default.")
    config = analysis._analysis_config(config)
    prepared = {key: deepcopy(dict(source.get("prepared_payload") or {}).get(key))
                for key in ("input_records", "original_points", "current_plan")}
    if not all(prepared.values()) or analysis._provider_for_records(prepared["input_records"]) != "amap":
        raise ValueError("Full correction requires the saved China workbook inputs.")
    if any(str(row.get("country") or "").lower() != "china" for row in prepared["input_records"]):
        raise ValueError("Full AMap correction does not accept mixed-market inputs.")
    current = prepared["current_plan"]
    if current.get("service_direction") != config["service_direction"]:
        raise ValueError("Saved workbook and result directions disagree.")
    groups: dict[str, list[dict[str, Any]]] = {}
    for stop in current.get("stops") or []:
        route_id = str(stop.get("route_id") or "").strip()
        riders = reviews._number(stop.get("passenger_count"))
        if (not route_id or str(stop["route_id"]) != route_id or riders is None
                or riders != int(riders)):
            raise ValueError("A saved stop lacks its route identity or passenger count.")
        groups.setdefault(route_id, []).append(stop)
    if not 1 <= len(groups) <= MAX_FULL_ROUTES:
        raise ValueError("Full correction requires 1-200 saved routes.")
    for stops in groups.values():
        sequences = [stop.get("stop_sequence") for stop in stops]
        if any(not isinstance(value, int) or isinstance(value, bool) for value in sequences) or len(set(sequences)) != len(stops):
            raise ValueError("Saved stop ordering is incomplete or ambiguous.")
        ordered = sorted(stops, key=lambda stop: stop["stop_sequence"])
        depot_indexes = [i for i, stop in enumerate(ordered) if stop.get("is_depot") is True]
        expected_depot = len(ordered) - 1 if config["service_direction"] == "To School" else 0
        if depot_indexes != [expected_depot] or len(ordered) < 2:
            raise ValueError("Saved school placement is inconsistent with service direction.")
    route_rows = list(saved.get("routes") or [])
    previous = {str(route.get("route_id") or ""): route for route in route_rows}
    if len(previous) != len(route_rows):
        raise ValueError("Saved route identities are ambiguous.")
    # Include input routes missing from a partial historical result, not just its successful rows.
    adapter_source = {**source, "result": {**saved, "parameters": config,
        "routes": [previous.get(route_id, {"route_id": route_id}) for route_id in sorted(groups)]}}
    routes = reviews.source_route_scopes(adapter_source)
    before = {"summary": deepcopy(saved.get("summary") or {}),
              "operational_conclusion": deepcopy(saved.get("operational_conclusion") or {}),
              "stops": [{key: deepcopy(row.get(key)) for key in
                        ("stop_key", "address", "riders", "operational_category")}
                        for row in saved.get("stops") or []]}
    full_input = {"prepared_payload": prepared, "analysis_config": config}
    return {
        "review_version": reviews.REVIEW_VERSION, "evidence_version": EVIDENCE_VERSION, "mode": MODE,
        "source_job_id": str(source["job_id"]), "owner_email": str(source.get("owner_email") or "").strip().lower(),
        "requested_by": requested_by.strip().lower(), "request_key": request_key,
        "source_finished_at": source.get("finished_at"), "source_result_digest": reviews._digest(source["result"]),
        "input_digest": reviews._digest(routes), "routes": routes, "provider_call_limit": provider_call_limit,
        "full_input": full_input, "full_input_digest": reviews._digest(full_input), "before_analysis": before,
        "scope_summary": {"route_count": len(groups),
                          "address_count": len({analysis._stop_key(stop) for stop in current["stops"] if not stop.get("is_depot")})},
    }


class FullReviewProvider(analysis.FreshRouteProvider):
    def __init__(self, provider: Any, *, read_snapshot: Callable, write_snapshot: Callable,
                 check_active: Callable[[], None]):
        self.inner = provider
        self.provider, self.state = provider.provider, provider.state
        self.read_snapshot, self.write_snapshot, self.check_active = read_snapshot, write_snapshot, check_active

    def route(self, points: list[dict[str, Any]], *, reference_legs=None) -> dict[str, Any]:
        self.check_active()
        self.state.pop("last_route_evidence", None)
        key = reviews._digest({"version": EVIDENCE_VERSION, "points": points, "references": reference_legs or []})
        snapshot = self.read_snapshot(key)
        if snapshot:
            try:
                captured_at = snapshot.get("measurement_started_at") or snapshot.get("called_at")
                age = (datetime.now(timezone.utc) - datetime.fromisoformat(str(captured_at).replace("Z", "+00:00"))).total_seconds()
            except (TypeError, ValueError):
                age = -1
            if 0 <= age <= CACHE_MAX_AGE_SECONDS and _verified(snapshot, len(points)):
                self.state["cache_hits"] = int(self.state.get("cache_hits", 0)) + 1
                snapshot.update(api_calls=0, in_run_reuse=True)
                self.state["last_route_evidence"] = deepcopy(snapshot)
                return snapshot
        try:
            value = self.inner.route(points, reference_legs=reference_legs)
        except reviews.ReviewClaimLost:
            raise
        except Exception as exc:
            raise RuntimeError(f"Route measurement failed ({type(exc).__name__}).") from None
        if not _verified(value, len(points)):
            self.state["last_route_evidence"] = deepcopy(value)
            raise RuntimeError("Full correction requires verified adjacent-leg measurements.")
        if not self.write_snapshot(key, value):
            raise reviews.ReviewClaimLost("Review stopped before its measurement snapshot was saved.")
        self.state["last_route_evidence"] = deepcopy(value)
        return value


def _verified(value: dict[str, Any], point_count: int) -> bool:
    if not (value.get("evidence_version") == EVIDENCE_VERSION and value.get("status") == "verified"
            and value.get("provider") == "amap" and not value.get("issues")
            and value.get("complete") is True and len(value.get("legs") or []) == point_count - 1
            and len(value.get("leg_durations_s") or []) == point_count - 1
            and reviews._number(value.get("duration_s")) is not None
            and reviews._number(value.get("distance_m")) is not None):
        return False
    durations = [reviews._number(leg.get("duration_s")) for leg in value["legs"]]
    distances = [reviews._number(leg.get("distance_m")) for leg in value["legs"]]
    return (None not in durations and None not in distances
            and durations == value["leg_durations_s"]
            and math.isclose(sum(durations), reviews._number(value["duration_s"]), abs_tol=0.01)
            and math.isclose(sum(distances), reviews._number(value["distance_m"]), abs_tol=0.01))


def _envelope(request: dict[str, Any], result: dict[str, Any], api_calls: int) -> dict[str, Any]:
    return {"review_version": reviews.REVIEW_VERSION, "source_job_id": request["source_job_id"],
            "source_result_digest": request["source_result_digest"], "full_input_digest": request["full_input_digest"],
            "mode": MODE, "scope": "full_direct_school_result", "status": "running", "routes": [],
            "provider_api_calls": api_calls, "analysis_result": deepcopy(result),
            "comparison_basis": "fresh_traffic_not_controlled_before_after",
            "student_classification_recomputed": False, "time_window_revalidated": False}


def run_full_review(store: Any, record: dict[str, Any], token: str, *,
                    provider_factory: Callable, checkpoint: Callable) -> dict[str, Any]:
    request, review_id = record["request"], record["review_id"]
    if (request.get("review_version") != reviews.REVIEW_VERSION or request.get("evidence_version") != EVIDENCE_VERSION
            or reviews._digest(request.get("full_input")) != request.get("full_input_digest")
            or reviews._digest(request["routes"]) != request.get("input_digest")):
        raise ValueError("Full correction input or contract changed.")
    used = int(record.get("api_calls") or 0)
    limit = request["provider_call_limit"]
    if not 0 <= used <= limit:
        raise ValueError("Invalid persisted provider usage.")

    def check_active():
        if not store.route_measurement_claim_active(review_id, token):
            raise reviews.ReviewClaimLost("Review is no longer running.")

    def factory(name, **kwargs):
        check_active()
        provider = provider_factory(name, **kwargs)
        provider.state = reviews.ReviewCallState({**provider.state, "api_calls": used},
            lambda count: store.reserve_route_measurement_calls(review_id, token, count))
        return FullReviewProvider(provider,
            read_snapshot=lambda key: store.get_review_measurement_snapshot(review_id, token, key),
            write_snapshot=lambda key, value: store.save_review_measurement_snapshot(review_id, token, key, value),
            check_active=check_active)

    config = {**request["full_input"]["analysis_config"], "provider_call_limit": limit}
    # Derived classifications are rerun; only fresh, review-owned road measurements survive a pause.
    result = analysis.run_direct_school_analysis(deepcopy(request["full_input"]["prepared_payload"]), config,
        run_seed=request["request_key"], provider_factory=factory, check_canceled=check_active,
        checkpoint=lambda partial: checkpoint(_envelope(request, partial, int(partial["progress"]["provider_api_calls"]))))
    check_active()
    output = _envelope(request, result, int(result["summary"]["provider_api_calls"]))
    current = {route["route_id"]: route for route in result["routes"]}
    for scope in request["routes"]:
        measured = current.get(scope["route_id"], {})
        certified = measured.get("status") == "resolved" and _verified(dict(measured.get("route_evidence") or {}), len(scope["points"]))
        evidence = dict(measured.get("route_evidence") or {})
        after = {"drive_duration_s": reviews._number(evidence.get("duration_s")) if certified else None,
                 "distance_m": reviews._number(evidence.get("distance_m")) if certified else None,
                 "total_duration_s": reviews._number(evidence["duration_s"]) + scope["stop_service_time_s"] if certified else None,
                 "captured_at": measured.get("provider_called_at")}
        output["routes"].append({"route_key": scope["route_key"], "route_id": scope["route_id"],
            "points": scope["points"], "status": "verified" if certified else "needs_review", "before": scope["before"],
            "after": after, "route_evidence": measured.get("route_evidence"),
            "delta": {key: after[key] - scope["before"][key] if after[key] is not None and reviews._number(scope["before"].get(key)) is not None else None
                      for key in ("drive_duration_s", "distance_m", "total_duration_s")}})
    unknown = dict(result["operational_conclusion"].get("data_review") or {})
    output["student_classification_recomputed"] = True
    output["classification_complete"] = not unknown.get("address_count") and not result["summary"]["failed_count"]
    output["time_window_revalidated"] = len(result["route_window_analysis"]) == len(request["routes"])
    output["time_window_revalidated"] &= not result["summary"]["route_window_data_review_count"]
    output["status"] = "complete" if (result["status"] == "complete" and output["classification_complete"]
        and output["time_window_revalidated"] and all(row["status"] == "verified" for row in output["routes"])) else "partial"
    before_summary = request["before_analysis"]["summary"]
    keys = ("direct_over_limit_rider_count", "route_only_over_limit_rider_count", "additional_removal_rider_count",
            "routes_over_window_final_count", "route_window_data_review_count")
    output["conclusion_changes"] = {key: {"before": reviews._number(before_summary.get(key)), "after": result["summary"].get(key),
        "delta": result["summary"][key] - reviews._number(before_summary[key]) if reviews._number(before_summary.get(key)) is not None else None} for key in keys}
    before_stops = {row["stop_key"]: row for row in request["before_analysis"]["stops"]}
    output["classification_changes"] = [{"stop_key": row["stop_key"], "address": row["address"], "riders": row["riders"],
        "before": before_stops.get(row["stop_key"], {}).get("operational_category"), "after": row["operational_category"]}
        for row in result["stops"] if before_stops.get(row["stop_key"], {}).get("operational_category") != row["operational_category"]]
    return output


def validate_full_result(request: dict[str, Any], result: dict[str, Any], *, terminal: bool) -> None:
    if result.get("status") == "failed":
        return
    if result.get("scope") != "full_direct_school_result" or result.get("full_input_digest") != request["full_input_digest"]:
        raise ValueError("Full correction result does not match its input snapshot.")
    analysis_result = dict(result.get("analysis_result") or {})
    expected_config = {**request["full_input"]["analysis_config"], "provider_call_limit": request["provider_call_limit"]}
    if analysis_result.get("parameters") != expected_config:
        raise ValueError("Full correction must retain the original analysis parameters.")
    if terminal and result.get("status") in {"complete", "partial"}:
        route_ids = sorted(scope["route_id"] for scope in request["routes"])
        if sorted(row.get("route_id") for row in analysis_result.get("routes") or []) != route_ids:
            raise ValueError("Full correction must retain every original route.")
        expected_stops = {analysis._stop_key(stop) for stop in request["full_input"]["prepared_payload"]["current_plan"]["stops"] if not stop.get("is_depot")}
        if {row.get("stop_key") for row in analysis_result.get("stops") or []} != expected_stops:
            raise ValueError("Full correction must retain every original service address.")
        expected_occurrences = Counter((analysis._stop_key(stop), stop["route_id"], stop["stop_sequence"],
                                        analysis._safe_int(stop["passenger_count"]))
            for stop in request["full_input"]["prepared_payload"]["current_plan"]["stops"] if not stop.get("is_depot"))
        actual_occurrences = Counter((row["stop_key"], item.get("route_id"), item.get("stop_sequence"), item.get("passenger_count"))
            for row in analysis_result["stops"] for item in row.get("occurrences") or [])
        if actual_occurrences != expected_occurrences or any(
                row.get("riders") != sum(item["passenger_count"] for item in row.get("occurrences") or [])
                for row in analysis_result["stops"]):
            raise ValueError("Full correction must retain each route occurrence and its student count.")


def build_full_review_workbook(record: dict[str, Any]) -> bytes:
    result = dict(record.get("result") or {})
    if (record.get("status") not in {"succeeded", "needs_review"}
            or result.get("scope") != "full_direct_school_result" or not result.get("analysis_result")):
        raise ValueError("A completed full Direct-to-School correction is required for this report.")
    content = analysis.build_direct_school_workbook({"job_id": record["review_id"],
        "metadata": {"job_name": f"Measurement correction of {record['source_job_id']}"},
        "result": result["analysis_result"]})
    workbook = load_workbook(BytesIO(content))
    labels = {"direct_over_limit_rider_count": "Students exceeding the direct-trip limit",
              "route_only_over_limit_rider_count": "Students exceeding the limit only on the shared route",
              "additional_removal_rider_count": "Additional students suggested for removal",
              "routes_over_window_final_count": "Routes still exceeding the window after removal",
              "route_window_data_review_count": "Routes awaiting measurement review"}
    changes_by_key = result.get("conclusion_changes", {})
    rows = [[label, changes_by_key[key].get("before"), changes_by_key[key].get("after"), changes_by_key[key].get("delta")]
            for key, label in labels.items() if key in changes_by_key]
    analysis._write_readable_table(workbook.create_sheet("Correction Comparison"),
        "Original result and corrected result",
        "Fresh traffic samples are not a controlled before/after experiment. Blank original values mean unavailable evidence.",
        ["Measure", "Original", "Corrected", "Change"], rows, [58, 18, 18, 18])
    changes = [[row["address"], row["riders"],
                analysis._operational_category_label(row["before"]) if row["before"] is not None else None,
                analysis._operational_category_label(row["after"])] for row in result.get("classification_changes") or []]
    analysis._write_readable_table(workbook.create_sheet("Classification Changes"),
        "Addresses with a changed classification", "All service addresses remain in the main classification worksheet.",
        ["Address", "Students", "Original classification", "Corrected classification"], changes, [58, 14, 45, 45])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
