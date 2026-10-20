"""Immutable morning/afternoon pairs and native return-route measurements."""
from collections import defaultdict
from copy import deepcopy
from datetime import datetime, timedelta, timezone
import hashlib
from io import BytesIO
import json
import math
from zoneinfo import ZoneInfo

VERSION = 1
LIMIT_S = 3600
MAX_DELAY_S = 15 * 60


def return_limit_minutes(value=60):
    if isinstance(value, bool):
        raise ValueError("Return duration limit must be between 1 and 240 minutes.")
    try:
        minutes = float(value)
    except (TypeError, ValueError) as exc:
        raise ValueError("Return duration limit must be between 1 and 240 minutes.") from exc
    if not math.isfinite(minutes) or not 1 <= minutes <= 240:
        raise ValueError("Return duration limit must be between 1 and 240 minutes.")
    return minutes


def under_return_limit(route, threshold_s):
    duration = route.get("total_duration_s")
    return (route.get("status") == "resolved" and isinstance(duration, (int, float))
            and not isinstance(duration, bool) and math.isfinite(duration)
            and 0 <= duration < threshold_s)


def fingerprint(value):
    return hashlib.sha256(json.dumps(value, sort_keys=True, ensure_ascii=False,
                                     allow_nan=False).encode()).hexdigest()


def reverse_payload(prepared):
    result = deepcopy(prepared)
    plan = result["current_plan"]
    if plan.get("service_direction") != "To School":
        raise ValueError("Paired return requires a To School plan.")
    groups = defaultdict(list)
    for stop in plan["stops"]:
        groups[str(stop["route_id"])].append(stop)
    if not groups:
        raise ValueError("The plan has no routes.")
    reversed_stops = []
    for route_id, stops in groups.items():
        stops.sort(key=lambda stop: int(stop["stop_sequence"]))
        sequences = [int(stop["stop_sequence"]) for stop in stops]
        if len(set(sequences)) != len(stops):
            raise ValueError(f"Duplicate stop sequence: {route_id}")
        if len(stops) < 2 or not stops[-1].get("is_depot") or any(s.get("is_depot") for s in stops[:-1]):
            raise ValueError(f"Route must end at one school: {route_id}")
        for sequence, stop in enumerate(reversed(stops)):
            stop["source_stop_sequence"] = stop["stop_sequence"]
            stop["stop_sequence"] = sequence
            reversed_stops.append(stop)
    plan["stops"] = reversed_stops
    plan["service_direction"] = "From School"
    # Morning geometry/timing and solver node lists are not return evidence.
    plan.pop("routes", None)
    result["service_direction"] = "From School"
    return result


def create_pair(service, payload, prepared, metadata, owner, morning_at, *, morning_config=None):
    metadata = deepcopy(metadata)
    metadata["analysis_config"]["return_duration_limit_minutes"] = return_limit_minutes(
        payload.get("return_duration_limit_minutes", 60))
    submission_id = str(payload.get("submission_id") or fingerprint([
        prepared["current_plan"], metadata["analysis_config"], morning_at,
        payload.get("return_scheduled_time", "15:40"), metadata.get("job_name")])).strip()
    if not submission_id or len(submission_id) > 128:
        raise ValueError("A submission identity is required for paired schedules.")
    afternoon_at, label = service._direct_school_schedule_trigger(
        scheduled_date=payload.get("scheduled_date"),
        scheduled_time=payload.get("return_scheduled_time") or "15:40",
        current_plan=prepared["current_plan"],
    )
    am = datetime.fromisoformat(morning_at)
    pm = datetime.fromisoformat(afternoon_at)
    if pm <= am or pm.hour < 12:
        raise ValueError("Return measurement must be in the afternoon after the morning run.")
    records = prepared.get("input_records") or []
    if service.infer_traffic_location(records)[0].strip().upper() not in {"CHINA", "CN"}:
        raise ValueError("Paired return schedules are available for China only.")
    reversed_prepared = reverse_payload(prepared)
    source_version = fingerprint(prepared["current_plan"])
    pair_id = fingerprint([owner.strip().lower(), service.JOB_QUEUE_SCOPE, submission_id,
                           payload.get("scheduled_date")])[:24]
    request_values = [source_version, metadata["analysis_config"], morning_at,
                      afternoon_at, metadata.get("job_name")]
    if morning_config is not None:
        request_values.append(morning_config)
    request_hash = fingerprint(request_values)
    ids = {"am": pair_id + "-am", "pm": pair_id + "-pm"}
    pair = []
    for period, snapshot, scheduled_at in (("am", prepared, morning_at), ("pm", reversed_prepared, afternoon_at)):
        meta = deepcopy(metadata)
        config = deepcopy(morning_config or {})
        meta.update(pair_id=pair_id, period=period, paired_job_ids=ids,
                    pair_request_hash=request_hash, source_plan_version=source_version,
                    scheduled_start_at=scheduled_at, scheduled_time=scheduled_at[11:16],
                    schedule_expires_at=(datetime.fromisoformat(scheduled_at) + timedelta(seconds=MAX_DELAY_S)).isoformat())
        if period == "pm":
            if morning_config is not None:
                config.update(service_direction="From School", final_time_validation_mode="legacy")
                meta["planner_config"] = config
                meta["return_dependency_status"] = "waiting_for_morning"
            meta["measurement_kind"] = "pm_return"
            meta["analysis_config"].update(service_direction="From School", final_time_validation_mode="legacy",
                                           from_school_departure_time=pm.strftime("%H:%M"))
            meta["analysis_signature"] = service._direct_school_signature(snapshot["current_plan"], meta["analysis_config"])
            meta["scheduled_trigger_label"] = pm.strftime("%H:%M") + " return measurement"
        stored_snapshot = {} if period == "pm" and morning_config is not None else snapshot
        record = service.JOB_STORE.create_job(config, stored_snapshot, metadata=meta, owner_email=owner,
            status="scheduled", scheduled_start_at=scheduled_at,
            scheduled_trigger_label=meta.get("scheduled_trigger_label") or label, persist=False)
        record["job_id"] = ids[period]
        pair.append(record)
    saved = service._runtime_sqlite_store().insert_job_pair(pair)
    from runtime_store_sqlite import job_summary
    return [job_summary(record) for record in saved]


def create_audit_pair(service, payload, prepared, metadata, owner, morning_at, config):
    meta = deepcopy(metadata)
    meta["job_kind"] = "route_audit"
    meta["analysis_config"] = {
        "service_direction": "To School", "final_time_validation_mode": "legacy",
        "stop_service_minutes": config.get("stop_service_minutes", 1),
        "far_duration_minutes": 60,
    }
    return create_pair(service, payload, prepared, meta, owner, morning_at, morning_config=config)


def ensure_return_due(metadata, now=None):
    now = now or datetime.now(timezone.utc)
    ensure_current_window(metadata, now)
    if metadata.get("measurement_kind") == "pm_return":
        scheduled = metadata.get("scheduled_start_at")
        if scheduled and now < datetime.fromisoformat(scheduled):
            raise ValueError("Return measurement cannot run before its scheduled traffic window.")


def ensure_current_window(metadata, now=None):
    expires = metadata.get("schedule_expires_at")
    if expires and (now or datetime.now(timezone.utc)) > datetime.fromisoformat(expires):
        raise ValueError("Scheduled measurement window expired. Create a new measurement; past traffic cannot be replayed.")


def _number(value):
    if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value) or value < 0:
        raise ValueError("Incomplete native route measurement.")
    return float(value)


def route_timings(ordered, evidence, dwell_s, departure, threshold_s=LIMIT_S):
    legs = [_number(value) for value in evidence.get("leg_durations_s", [])]
    if evidence.get("complete") is not True or len(legs) != len(ordered) - 1:
        raise ValueError("Incomplete native route measurement.")
    drive_s = _number(evidence.get("duration_s"))
    if abs(sum(legs) - drive_s) > max(1, drive_s * .001):
        raise ValueError("Native route total and stop timings disagree.")
    dwell_s = _number(dwell_s)
    elapsed = 0.0
    stops = []
    for index, original in enumerate(ordered):
        if index:
            elapsed += legs[index - 1]
        stop = deepcopy(original)
        stop.update(ride_duration_s=elapsed, arrival_at=(departure + timedelta(seconds=elapsed)).isoformat())
        service_s = 0.0 if stop.get("is_depot") else dwell_s
        stop["departure_at"] = (departure + timedelta(seconds=elapsed + service_s)).isoformat()
        stops.append(stop)
        elapsed += service_s
    return {"drive_duration_s": drive_s, "dwell_duration_s": elapsed - drive_s,
            "total_duration_s": elapsed, "under_one_hour": elapsed < LIMIT_S,
            "under_duration_limit": elapsed < threshold_s, "stops": stops}


def run_return_measurement(prepared, config, *, checkpoint=None, check_canceled=None, provider_factory=None, now=None, provider=None):
    import direct_school_analysis as analysis
    threshold_s = return_limit_minutes(config.get("return_duration_limit_minutes", 60)) * 60
    if config.get("final_time_validation_mode", "legacy") != "legacy":
        raise ValueError("Return measurements use the existing map provider only.")
    plan = prepared["current_plan"]
    if plan.get("service_direction") != "From School":
        raise ValueError("Return measurement requires a reversed From School snapshot.")
    input_records = prepared["input_records"]
    lookup = analysis._point_lookup(input_records, prepared["original_points"])
    provider_name = analysis._provider_for_records(input_records)
    if provider_name != "amap":
        raise ValueError("Return measurements require the China map provider.")
    clock = now or (lambda: datetime.now(timezone.utc))
    provider = provider or (provider_factory or analysis.FreshRouteProvider)(provider_name,
        departure_time=None, api_call_limit=int(config.get("provider_call_limit", 500)))
    initial_calls = provider.state.get("api_calls", 0)
    initial_hits = provider.state.get("cache_hits", 0)
    groups = defaultdict(list)
    for stop in plan["stops"]:
        groups[str(stop["route_id"])].append(stop)
    result = {"analysis_kind": "pm_return", "analysis_version": VERSION, "service_direction": "From School",
              "provider": provider_name, "status": "running", "started_at": clock().isoformat(),
              "parameters": deepcopy(config), "routes": [], "summary": {}, "progress": {},
              "timing_basis": "current_traffic_snapshot", "threshold_s": threshold_s}
    osrm_cache = {}
    for route_id, unordered in groups.items():
        if check_canceled:
            check_canceled()
        ordered = sorted(unordered, key=lambda stop: int(stop["stop_sequence"]))
        route = {"route_id": route_id, "riders": sum(int(s.get("passenger_count") or 0) for s in ordered),
                 "stop_count": len(ordered) - 1, "status": "failed", "under_one_hour": False, "under_duration_limit": False,
                 "stops": deepcopy(ordered)}
        route.update(deepcopy(ordered[0].get("source_route") or {}))
        result["routes"].append(route)
        try:
            points = [stop.get("source_point") or analysis._lookup_point(stop, lookup) for stop in ordered]
            if any(point is None for point in points):
                raise ValueError("One or more return stops could not be geocoded.")
            for stop, point in zip(route["stops"], points):
                coords = analysis._point_coordinates(point)
                if coords:
                    stop.update(lat=coords[0], lng=coords[1])
            refs = [analysis._osrm_leg(a, b, osrm_cache) for a, b in zip(points, points[1:])]
            measured_at = clock()
            native = analysis._route_with_stop_boundaries(provider, points, reference_legs=refs)
            timings = route_timings(route["stops"], native, float(config["stop_service_minutes"]) * 60, measured_at, threshold_s)
            for stop in timings["stops"]:
                stop["over_limit"] = not stop.get("is_depot") and stop["ride_duration_s"] > float(config.get("far_duration_minutes", 45)) * 60
            timings["over_limit_riders"] = sum(int(s.get("passenger_count") or 0) for s in timings["stops"] if s["over_limit"])
            route.update(timings, status="resolved", measured_at=native.get("called_at") or measured_at.isoformat(),
                         departure_at=measured_at.isoformat(), distance_m=_number(native.get("distance_m")),
                         geometry=native.get("geometry") or [], geometry_segments=native.get("geometry_segments"),
                         provider_route_source=native.get("source") or provider_name, route_evidence=native)
        except Exception as exc:
            route.update(error=str(exc), under_one_hour=False)
        result["summary"] = {"route_count": len(groups), "completed_route_count": sum(r["status"] == "resolved" for r in result["routes"]),
            "under_one_hour_count": sum(r["under_one_hour"] for r in result["routes"]),
            "under_duration_limit_count": sum(under_return_limit(r, threshold_s) for r in result["routes"]),
            "rider_count": sum(int(s.get("passenger_count") or 0) for s in plan["stops"])}
        result["progress"] = {"completed": len(result["routes"]), "total": len(groups),
                              "provider_api_calls": provider.state.get("api_calls", 0) - initial_calls,
                              "in_run_reuse_count": provider.state.get("cache_hits", 0) - initial_hits}
        if checkpoint:
            checkpoint(deepcopy(result))
    result["status"] = "complete" if all(r["status"] == "resolved" for r in result["routes"]) else "partial"
    result["completed_at"] = clock().isoformat()
    return result


def build_return_workbook(record, duration_limit_minutes=None):
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill
    result = record.get("result") or {}
    saved_config = (record.get("metadata") or {}).get("analysis_config") or {}
    minutes = return_limit_minutes(duration_limit_minutes if duration_limit_minutes is not None else
        (result.get("parameters") or saved_config).get("return_duration_limit_minutes", result.get("threshold_s", LIMIT_S) / 60))
    threshold_s = minutes * 60
    limit_label = f"Under {minutes:g} min / 低于{minutes:g}分钟"
    workbook = Workbook()
    if result.get("scenarios") is not None:
        from openpyxl import load_workbook
        overview = workbook.active
        overview.title = "Scenarios"
        overview.append(["Source morning job", result.get("source_job_id")])
        overview.append(["Source version", result.get("source_version")])
        overview.append(["Scheduled", record.get("scheduled_start_at"), "Measured", result.get("started_at")])
        overview.append(["Scenario", "Morning status", "Return status", "Routes", limit_label, "Error"])
        labels = {"current_plan": "Current", "time_constrained": "Strict", "exception_preserving": "Protected"}
        for key, scenario in result["scenarios"].items():
            label = labels[key]
            overview.append([label, scenario.get("source_status"), scenario.get("status"),
                             scenario.get("summary", {}).get("route_count"),
                             sum(under_return_limit(r, threshold_s) for r in scenario.get("routes", [])), scenario.get("error")])
            child = load_workbook(BytesIO(build_return_workbook({**record, "result": scenario}, minutes)))
            for original in child:
                sheet = workbook.create_sheet(f"{label} {original.title}")
                for row in original:
                    for cell in row:
                        target = sheet.cell(cell.row, cell.column, cell.value)
                        from copy import copy
                        target._style = copy(cell._style)
                sheet.freeze_panes = original.freeze_panes
                for column, dimension in original.column_dimensions.items():
                    sheet.column_dimensions[column].width = dimension.width
        overview.freeze_panes = "A5"
        for column in overview.columns:
            overview.column_dimensions[column[0].column_letter].width = 28
        stream = BytesIO()
        workbook.save(stream)
        return stream.getvalue()
    overview = workbook.active
    overview.title = "Return routes"
    overview.append(["Direction / 方向", "From School / 下午返程"])
    overview.append(["Scheduled / 计划执行", record.get("scheduled_start_at")])
    overview.append(["Measured / 实际开始", result.get("started_at")])
    overview.append(["Basis / 口径", "Current traffic snapshot; driving + all stop dwell / 实时快照：行车加全部停站时间"])
    overview.append(["Route / 路线", "Status / 状态", "Riders / 人数", "Driving (s) / 行车秒数", "Dwell (s) / 停站秒数",
                     "Total (s) / 总秒数", limit_label, "Measured at / 测算时间", "Provider / 来源"])
    stops = workbook.create_sheet("Return stops")
    stops.append(["Route / 路线", "Return order / 返程序号", "Source order / 原序号", "Address / 地址", "Riders / 人数",
                  "Ride (s) / 乘车秒数", "Arrival / 到达时间", "Departure / 离站时间", "Over trip limit / 乘车超时"])
    for route in result.get("routes", []):
        overview.append([route["route_id"], route["status"], route["riders"], route.get("drive_duration_s"),
                         route.get("dwell_duration_s"), route.get("total_duration_s"),
                         "Yes / 是" if under_return_limit(route, threshold_s) else "No / 否" if route["status"] == "resolved" else "Unavailable / 未完成",
                         route.get("measured_at"), route.get("provider_route_source")])
        if under_return_limit(route, threshold_s):
            for cell in overview[overview.max_row]:
                cell.fill = PatternFill("solid", fgColor="DCFCE7")
        for stop in route["stops"]:
            stops.append([route["route_id"], stop["stop_sequence"], stop.get("source_stop_sequence"), stop.get("address"),
                          stop.get("passenger_count"), stop.get("ride_duration_s"), stop.get("arrival_at"), stop.get("departure_at"), stop.get("over_limit")])
    for sheet in workbook:
        sheet.freeze_panes = "A6" if sheet == overview else "A2"
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = min(55, max(18, max(len(str(c.value or "")) for c in column) + 2))
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()
