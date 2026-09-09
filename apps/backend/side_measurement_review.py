"""Native Fleet/Insert corrections with immutable sources and bounded road snapshots."""
from __future__ import annotations

from copy import deepcopy
from typing import Any, Callable

try:
    from . import measurement_reviews as reviews
    from .full_measurement_review import FullReviewProvider, _verified
except ImportError:
    import measurement_reviews as reviews
    from full_measurement_review import FullReviewProvider, _verified

MODES = {"fleet_planner": "full_fleet", "route_insert_advisor": "full_insert"}
PLAN_FIELDS = ("global_plan_result", "route_preview_result")
CN = {"CHINA", "CN", "中国", "中华人民共和国"}
MAX_ROUTES = 200


def _payload(source: dict, tool: str) -> dict:
    if tool == "fleet_planner":
        return {key: deepcopy(source.get(key)) for key in (*PLAN_FIELDS, "geocode_result", "scenario")}
    return {"route_insert_result": deepcopy(source.get("route_insert_result"))}


def _identity(point: dict) -> tuple:
    return (point.get("address"), point.get("lat"), point.get("lng"),
            point.get("is_depot"), point.get("passenger_count"))


def _points_valid(points: list[dict], direction: str) -> bool:
    if len(points) < 2 or direction not in {"to_school", "from_school"}:
        return False
    school = len(points) - 1 if direction == "to_school" else 0
    for index, point in enumerate(points):
        lat, lng = reviews._number(point.get("lat")), reviews._number(point.get("lng"))
        count = reviews._number(point.get("passenger_count"))
        if (lat is None or lng is None or not 0 < lat <= 90 or not 0 < lng <= 180
                or count is None or not count.is_integer() or not point.get("address")
                or point.get("is_depot") is not (index == school)
                or str(point.get("country") or "").upper() not in CN
                or point.get("coordinate_system") not in {"WGS84", "GCJ02"}):
            return False
    return points[school]["passenger_count"] == 0


def _before(route: dict, dwell: Any) -> dict:
    evidence = dict(route.get("route_evidence") or {})
    gate = dict(route.get("final_route_traffic_gate") or {})
    drive = reviews._number(evidence.get("duration_s")) if evidence.get("status") == "verified" else reviews._number(gate.get("verified_drive_duration_s"))
    total = drive + dwell if drive is not None and dwell is not None else reviews._number(gate.get("verified_total_duration_s"))
    return {"drive_duration_s": drive, "total_duration_s": total,
            "distance_m": reviews._number(evidence.get("distance_m")) if evidence.get("status") == "verified" else reviews._number(gate.get("verified_distance_m")),
            "captured_at": evidence.get("called_at"),
            "planned_total_duration_s": reviews._number(route.get("duration_s")),
            "planned_distance_m": reviews._number(route.get("distance_m"))}


def _fleet_scopes(source: dict) -> list[dict]:
    scopes = []
    geocode_result = dict(source.get("geocode_result") or {})
    geocodes = [item for field in ("points", "demand_points")
                for item in (geocode_result.get(field) or []) if isinstance(item, dict)]
    if isinstance(geocode_result.get("school"), dict):
        geocodes.append(geocode_result["school"])
    for field in PLAN_FIELDS:
        plan = dict(source.get(field) or {})
        summary, school = dict(plan.get("summary") or {}), dict(plan.get("school") or {})
        direction = str(summary.get("service_direction") or "")
        target = reviews._number(summary.get("max_route_duration_minutes"))
        for route in plan.get("routes") or []:
            ordered = deepcopy(route.get("ordered_points") or [])
            points = []
            for index, point in enumerate(ordered):
                if not point.get("coordinate_system"):
                    matches = [item for item in geocodes if item.get("address") == point.get("address")
                               and item.get("lat") == point.get("lat") and item.get("lng") == point.get("lng")]
                    systems = {item.get("coordinate_system") for item in matches}
                    if len(systems) == 1 and systems <= {"WGS84", "GCJ02"}:
                        point = {**point, "coordinate_system": next(iter(systems))}
                depot = index == (len(ordered) - 1 if direction == "to_school" else 0)
                points.append({**point, "country": point.get("country") or school.get("country"),
                    "is_depot": depot, "passenger_count": point.get("student_count", point.get("passenger_count"))})
            dwell = reviews._number(route.get("stop_service_time_s"))
            scope = reviews._scope(f"{field}:{route.get('cluster_id', '')}",
                {**route, "route_id": route.get("cluster_id"), "vehicle_id": route.get("vehicle_id")},
                points, direction, dwell, target * 60 if target is not None else None,
                _before(route, dwell), route.get("leg_details"))
            endpoint = points[-1 if direction == "to_school" else 0] if points else {}
            if not _points_valid(points, direction) or any(endpoint.get(key) != school.get(key) for key in ("address", "lat", "lng")):
                scope["input_issues"].append("source_stop_coordinates_missing")
            if dwell is None or target is None or reviews._number(dict(route.get("selected_vehicle") or {}).get("student_capacity")) is None:
                scope["input_issues"].append("source_acceptance_inputs_missing")
            scope.update(plan_key=field, before_input=deepcopy(route))
            scopes.append(scope)
    return scopes


def _insert_scopes(source: dict) -> list[dict]:
    result = dict(source.get("route_insert_result") or {})
    scopes = []
    for scenario in result.get("scenarios") or []:
        plan = dict(scenario.get("selected_plan") or {})
        for route in plan.get("affected_routes") or []:
            saved = dict(route.get("measurement_inputs") or {})
            config = dict(saved.get("config") or {})
            direction = config.get("service_direction")
            direction = {"To School": "to_school", "From School": "from_school"}.get(direction, direction)
            for role in ("base", "selected"):
                evidence = route.get("base_route_evidence" if role == "base" else "route_evidence")
                dwell = reviews._number(saved.get(f"{role}_stop_service_time_s"))
                points = deepcopy(saved.get(f"{role}_points") or [])
                wrapped = {"route_id": route.get("route_id"), "route_evidence": evidence}
                scope = reviews._scope(f"{scenario.get('id', '')}:{route.get('route_id', '')}:{role}", wrapped,
                    points, str(direction or ""), dwell, saved.get("window_limit_s"), _before(wrapped, dwell))
                if saved.get("version") != 1 or not _points_valid(points, direction) or any(point.get("coordinate_system") != "WGS84" for point in points):
                    scope["input_issues"].append("source_stop_coordinates_missing")
                if (dwell is None or reviews._number(saved.get("inserted_stop_dwell_s")) is None
                        or reviews._number(saved.get("window_limit_s")) is None
                        or not config.get("time_window_start") or not config.get("time_window_end")
                        or reviews._number(route.get("capacity_limit")) is None):
                    scope["input_issues"].append("source_acceptance_inputs_missing")
                scope.update(plan_key=scenario.get("id"), role=role, before_input=deepcopy(route))
                scopes.append(scope)
    return scopes


def source_scopes(source: dict, tool: str) -> list[dict]:
    if tool not in MODES:
        raise ValueError("This tool has no native correction adapter.")
    return _fleet_scopes(source) if tool == "fleet_planner" else _insert_scopes(source)


def build_side_review_request(source: dict, *, tool_key: str, requested_by: str,
                              request_key: str, provider_call_limit: int) -> dict:
    if (tool_key not in MODES or not source.get("run_id") or not requested_by.strip()
            or not 1 <= len(request_key) <= 80 or source.get("status") in {"queued", "running", "cancelling"}):
        raise ValueError("A finished native source and requesting administrator are required.")
    if isinstance(provider_call_limit, bool) or not isinstance(provider_call_limit, int) or not 1 <= provider_call_limit <= 500:
        raise ValueError("Provider call limit must be an integer from 1 to 500.")
    scopes = source_scopes(source, tool_key)
    if not 1 <= len(scopes) <= MAX_ROUTES or len({row["route_key"] for row in scopes}) != len(scopes):
        raise ValueError("A full correction requires 1-200 uniquely identified saved route measurements.")
    if any(row["input_issues"] for row in scopes):
        raise ValueError("Saved coordinates, school roles, dwell, window or capacity inputs are incomplete. Re-prepare a new plan; do not infer historical inputs.")
    if tool_key == "route_insert_advisor":
        _validate_insert_inputs(source)
    payload = _payload(source, tool_key)
    return {"review_version": reviews.REVIEW_VERSION, "evidence_version": reviews.EVIDENCE_VERSION,
        "mode": MODES[tool_key], "source_job_id": None, "source_tool_key": tool_key,
        "source_run_id": source["run_id"], "owner_email": str(source.get("owner_email") or "").lower(),
        "requested_by": requested_by.strip().lower(), "request_key": request_key,
        "source_result_digest": reviews._digest(payload), "full_input": payload,
        "full_input_digest": reviews._digest(payload), "input_digest": reviews._digest(scopes),
        "routes": scopes, "provider_call_limit": provider_call_limit,
        "scope_summary": {"route_count": len(scopes), "plan_count": len({row["plan_key"] for row in scopes})}}


def side_risk_summary(source: dict, tool: str, email: str) -> dict:
    scopes = source_scopes(source, tool)
    risk = {"source_job_id": None, "source_tool_key": tool, "source_run_id": source.get("run_id"),
        "source_supported": bool(scopes), "status": "ready" if scopes else "unavailable",
        "routes": [{key: row[key] for key in ("route_key", "route_id", "risk_reasons", "input_issues")}
                   for row in scopes if row["risk_reasons"] or row["input_issues"]]}
    try:
        request = build_side_review_request(source, tool_key=tool, requested_by=email,
            request_key="availability-check", provider_call_limit=500)
        risk["full_review"] = {"mode": MODES[tool], "available": True, "scope_summary": request["scope_summary"]}
    except ValueError as exc:
        risk["full_review"] = {"mode": MODES[tool], "available": False, "reason": str(exc)}
    return risk


def _validate_insert_inputs(source: dict) -> None:
    try:
        from . import api_app as api
    except ImportError:
        import api_app as api
    for scenario in source["route_insert_result"].get("scenarios") or []:
        plan = scenario["selected_plan"]
        actions = plan.get("actions") or []
        rows = plan.get("affected_routes") or []
        if not actions or {a.get("route_id") for a in actions} != {r.get("route_id") for r in rows}:
            raise ValueError("Saved Insert actions must cover exactly its affected routes.")
        configs = [row["measurement_inputs"]["config"] for row in rows]
        if any(config != configs[0] for config in configs):
            raise ValueError("Saved Insert route settings are inconsistent.")
        for row in rows:
            saved = row["measurement_inputs"]
            relevant = [a for a in actions if a["route_id"] == row["route_id"]]
            base = saved["base_points"]
            counts = [reviews._number(row.get(key)) for key in ("capacity_before", "capacity_after", "base_stop_count", "selected_stop_count", "capacity_limit")]
            if (any(value is None or not value.is_integer() for value in counts) or counts[-1] <= 0
                    or [p.get("order") for p in base] != list(range(len(base)))
                    or api._insert_time_window_seconds(saved["config"]) != saved["window_limit_s"]):
                raise ValueError("Saved Insert capacities, order or time-window inputs are inconsistent.")
            inserted = [a for a in relevant if a.get("type") == "insert_stop"]
            if any(a.get("type") not in {"insert_stop", "walk_to_stop"} for a in relevant):
                raise ValueError("Unsupported saved Insert action.")
            sequence = api._insert_route_sequence(base, inserted)
            if [_identity(p) for p in sequence] != [_identity(p) for p in saved["selected_points"]]:
                raise ValueError("Saved selected route does not match its unchanged insertion actions.")
            if saved["selected_stop_service_time_s"] != saved["base_stop_service_time_s"] + len(inserted) * saved["inserted_stop_dwell_s"]:
                raise ValueError("Saved Insert dwell does not reconcile with its actions.")
            for action in relevant:
                if action["type"] == "walk_to_stop" and not any(p.get("order") == action.get("target_stop_order") and not p["is_depot"] for p in base):
                    raise ValueError("Saved walking target is missing or is the school.")


def run_side_review(store: Any, record: dict, token: str, *, provider_factory: Callable,
                    checkpoint: Callable) -> dict:
    request, review_id = record["request"], record["review_id"]
    if (request.get("mode") != MODES.get(request.get("source_tool_key"))
            or request.get("review_version") != reviews.REVIEW_VERSION
            or request.get("evidence_version") != reviews.EVIDENCE_VERSION
            or reviews._digest(request.get("full_input")) != request.get("full_input_digest")
            or reviews._digest(request.get("routes")) != request.get("input_digest")):
        raise ValueError("Native correction input or contract changed.")
    def active():
        if not store.route_measurement_claim_active(review_id, token):
            raise reviews.ReviewClaimLost("Review is no longer active.")
    active()
    inner = provider_factory("amap", departure_time=None, api_call_limit=request["provider_call_limit"])
    inner.state = reviews.ReviewCallState({**inner.state, "api_calls": int(record.get("api_calls") or 0)},
        lambda count: store.reserve_route_measurement_calls(review_id, token, count))
    provider = FullReviewProvider(inner, check_active=active,
        read_snapshot=lambda key: store.get_review_measurement_snapshot(review_id, token, key),
        write_snapshot=lambda key, value: store.save_review_measurement_snapshot(review_id, token, key, value))
    native = deepcopy(request["full_input"])
    output = {key: request[key] for key in ("review_version", "source_job_id", "source_tool_key", "source_run_id", "source_result_digest", "full_input_digest", "mode")}
    output.update(scope=request["mode"] + "_result", status="running", routes=[], native_result=native,
        provider_api_calls=int(inner.state["api_calls"]), comparison_basis="fresh_traffic_not_controlled_before_after",
        student_classification_recomputed=False, time_window_revalidated=False)
    measured: dict[str, dict] = {}
    if request["mode"] == "full_fleet":
        try:
            from . import backend_service as backend
        except ImportError:
            import backend_service as backend
        fleet = backend._client_module("demand_routing")
        for field in PLAN_FIELDS:
            plan = native.get(field)
            if not plan or not plan.get("routes"):
                continue
            for scope in [row for row in request["routes"] if row["plan_key"] == field]:
                route = next(row for row in plan["routes"] if row["cluster_id"] == scope["route_id"])
                route["ordered_points"] = deepcopy(scope["points"])
                route["stop_service_time_s"] = scope["stop_service_time_s"]
                active()
            backend._attach_fleet_route_measurements(plan, fleet, measurement_provider=provider, preserve_saved_dwell=True)
            active()
            for scope in [row for row in request["routes"] if row["plan_key"] == field]:
                route = next(row for row in plan["routes"] if row["cluster_id"] == scope["route_id"])
                measured[scope["route_key"]] = dict(route.get("route_evidence") or {})
            native[field] = backend._ensure_fleet_planner_map_data(plan)
            checkpoint({**output, "provider_api_calls": int(inner.state["api_calls"])})
    else:
        _run_insert(native, request, provider, measured, active)
    for scope in request["routes"]:
        evidence = measured.get(scope["route_key"], {})
        valid = _verified(evidence, len(scope["points"]))
        after = {"drive_duration_s": evidence.get("duration_s") if valid else None,
            "total_duration_s": evidence["duration_s"] + scope["stop_service_time_s"] if valid else None,
            "distance_m": evidence.get("distance_m") if valid else None, "captured_at": evidence.get("called_at")}
        output["routes"].append({"route_key": scope["route_key"], "route_id": scope["route_id"], "plan_key": scope["plan_key"],
            "points": scope["points"], "status": "verified" if valid else "needs_review", "before": scope["before"], "after": after,
            "route_evidence": evidence or None,
            "delta": {key: after[key] - scope["before"][key] if after[key] is not None and reviews._number(scope["before"].get(key)) is not None else None
                      for key in ("drive_duration_s", "total_duration_s", "distance_m")}})
    active()
    complete = all(row["status"] == "verified" for row in output["routes"])
    output.update(status="complete" if complete else "partial", time_window_revalidated=complete,
                  provider_api_calls=int(inner.state["api_calls"]))
    return output


def _run_insert(native: dict, request: dict, provider: Any, measured: dict, active: Callable) -> None:
    try:
        from . import api_app as api
    except ImportError:
        import api_app as api
    result = native["route_insert_result"]
    for scenario in result["scenarios"]:
        previous = scenario["selected_plan"]
        base_map = deepcopy(scenario.get("selected_map_data") or {})
        base_map.update(routes=[], stops=[])
        for row in previous["affected_routes"]:
            saved = row["measurement_inputs"]
            base_map["routes"].append({"id": row["route_id"], "load": row["capacity_before"], "bus_capacity": row["capacity_limit"],
                "stop_count": row["base_stop_count"], "stop_service_time_s": saved["base_stop_service_time_s"]})
            base_map["stops"].extend({**deepcopy(point), "route_id": row["route_id"]} for point in saved["base_points"])
        def measure(points, _country, _cache):
            active()
            try:
                evidence = provider.route(points)
            except reviews.ReviewClaimLost:
                raise
            except RuntimeError:
                evidence = deepcopy(provider.state.get("last_route_evidence") or {})
            valid = _verified(evidence, len(points))
            return {"geometry": evidence.get("geometry") or [], "display_geometry": evidence.get("geometry") or [],
                "display_geometry_source": evidence.get("source") or "amap_unavailable", "route_evidence": evidence,
                "duration_s": evidence.get("duration_s") if valid else None, "distance_m": evidence.get("distance_m") if valid else None,
                "leg_durations_s": evidence.get("leg_durations_s") or [] if valid else [],
                "leg_distances_m": evidence.get("leg_distances_m") or [] if valid else [], "provider_verified": valid}
        saved = previous["affected_routes"][0]["measurement_inputs"]
        config = {**saved["config"], "stop_service_minutes": saved["inserted_stop_dwell_s"] / 60}
        plan, map_data = api._insert_build_selected_plan(base_map, deepcopy(previous["actions"]),
            country="China", constraints={}, suggested_config=config, measurement_cache={}, measure_route=measure)
        scenario.update(selected_plan=plan, selected_map_data=map_data)
        by_route = {row["route_id"]: row for row in plan["affected_routes"]}
        for recommendation in scenario.get("recommendations") or []:
            selected = recommendation.get("selected")
            if not selected or selected.get("route_id") not in by_route:
                continue
            row = by_route[selected["route_id"]]
            selected.update(feasible=row["feasible"], measurement_scope="combined_selected_plan",
                base_route_duration_s=row["base_duration_s"], estimated_route_duration_s=row["selected_duration_s"],
                base_route_distance_m=row["base_distance_m"], estimated_route_distance_m=row["selected_distance_m"],
                delta_duration_s=row["delta_duration_s"], delta_distance_m=row["delta_distance_m"])
        for row in plan["affected_routes"]:
            measured[f"{scenario['id']}:{row['route_id']}:base"] = dict(row.get("base_route_evidence") or {})
            measured[f"{scenario['id']}:{row['route_id']}:selected"] = dict(row.get("route_evidence") or {})
        active()
    first = result["scenarios"][0]
    result.update(selected_plan=first["selected_plan"], selected_map_data=first["selected_map_data"], map_data=first["selected_map_data"],
                  recommendations=deepcopy(first.get("recommendations") or []))
    result["summary"].update({key: first["selected_plan"].get(key) for key in
        ("total_added_duration_s", "total_added_distance_m", "affected_route_count", "provider_verified_route_count")})


def validate_side_result(request: dict, result: dict, *, terminal: bool) -> None:
    if result.get("status") == "failed":
        return
    if result.get("scope") != request["mode"] + "_result" or result.get("full_input_digest") != request["full_input_digest"]:
        raise ValueError("Native correction does not match its saved input.")
    if terminal:
        actual = source_scopes(result.get("native_result") or {}, request["source_tool_key"])
        if [row["route_key"] for row in actual] != [row["route_key"] for row in request["routes"]]:
            raise ValueError("Native correction lost a saved plan or route.")
        for old, new in zip(request["routes"], actual):
            if ([_identity(p) for p in old["points"]] != [_identity(p) for p in new["points"]]
                    or [(p.get("coordinate_system"), p.get("country"), p.get("provider")) for p in old["points"]]
                    != [(p.get("coordinate_system"), p.get("country"), p.get("provider")) for p in new["points"]]
                    or new["stop_service_time_s"] != old["stop_service_time_s"]
                    or new["window_limit_s"] != old["window_limit_s"] or new["input_issues"]):
                raise ValueError("Native correction changed coordinates, riders, dwell or acceptance inputs.")
        comparisons = result.get("routes") or []
        if len(comparisons) != len(actual):
            raise ValueError("Native correction measurement coverage is incomplete.")
        for scope, comparison in zip(actual, comparisons):
            evidence = dict(scope.get("before_evidence") or {})
            valid = _verified(evidence, len(scope["points"]))
            after = {"drive_duration_s": evidence.get("duration_s") if valid else None,
                     "total_duration_s": evidence["duration_s"] + scope["stop_service_time_s"] if valid else None,
                     "distance_m": evidence.get("distance_m") if valid else None, "captured_at": evidence.get("called_at")}
            if (dict(comparison.get("route_evidence") or {}) != evidence or comparison.get("after") != after
                    or comparison.get("status") != ("verified" if valid else "needs_review")):
                raise ValueError("Native output and comparison must use the same road measurement.")


def corrected_native_record(record: dict) -> dict:
    result = dict(record.get("result") or {})
    if (record.get("status") not in {"succeeded", "needs_review"} or result.get("status") not in {"complete", "partial"}
            or result.get("scope") != record["request"]["mode"] + "_result" or not result.get("native_result")):
        raise ValueError("A finalized native correction is required.")
    return {**deepcopy(result["native_result"]), "run_id": record["source_run_id"],
            "measurement_review_id": record["review_id"]}


def build_side_review_workbook(record: dict, language: str = "en") -> bytes:
    from io import BytesIO
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    corrected_native_record(record)
    if language not in {"en", "zh", "ko"}:
        raise ValueError("Unsupported report language.")
    labels = {
        "en": ["Review Summary", "Route Measurements", "Stop Details", "Source", "Correction", "Executed", "API calls",
               "Plan", "Route", "Role", "Original min", "Corrected min", "Change min", "Original km", "Corrected km",
               "Students", "Measured at", "Status", "Order", "Address", "School", "Drive min", "Distance km",
               "Measured", "Needs review", "Base", "Selected", "Yes", "No",
               "Fresh traffic sample; stop order unchanged. This is not a controlled before/after experiment."],
        "zh": ["复核摘要", "线路测量", "站点明细", "原始记录", "修正记录", "执行时间", "API调用次数",
               "方案", "线路", "角色", "原时长(分)", "修正时长(分)", "变化(分)", "原距离(公里)", "修正距离(公里)",
               "学生人数", "测量时间", "状态", "顺序", "地址", "学校", "行驶时间(分)", "距离(公里)",
               "已测量", "待复核", "原线路", "选定线路", "是", "否",
               "本次为新的路况样本，站序不变；不是受控的修复前后实验。"],
        "ko": ["검토 요약", "노선 측정", "정류장 상세", "원본 기록", "수정 기록", "실행 시각", "API 호출",
               "계획", "노선", "역할", "기존 시간(분)", "수정 시간(분)", "변화(분)", "기존 거리(km)", "수정 거리(km)",
               "학생 수", "측정 시각", "상태", "순서", "주소", "학교", "주행 시간(분)", "거리(km)",
               "측정됨", "검토 필요", "기존 노선", "선택 노선", "예", "아니요",
               "새 교통 표본이며 정류장 순서는 동일합니다. 통제된 수정 전후 실험은 아닙니다."],
    }[language]
    wb = Workbook()
    summary = wb.active
    summary.title = labels[0]
    for key, value in ((labels[3], record["source_run_id"]), (labels[4], record["review_id"]),
                       (labels[5], record.get("started_at")), (labels[6], f"{record['api_calls']} / {record['request']['provider_call_limit']}")):
        summary.append([key, value])
    summary.append([labels[29]])
    summary.merge_cells("A5:B5")
    summary["A5"].alignment = Alignment(wrap_text=True, vertical="top")
    summary.row_dimensions[5].height = 45
    routes = wb.create_sheet(labels[1])
    routes.append([labels[i] for i in (7, 8, 9, 10, 11, 12, 13, 14, 15, 16, 17)])
    stops = wb.create_sheet(labels[2])
    stops.append([labels[i] for i in (7, 8, 9, 18, 19, 15, 20, 21, 22)])
    scope_by_key = {row["route_key"]: row for row in record["request"]["routes"]}
    def metric(value, divisor):
        return round(value / divisor, 2) if isinstance(value, (float, int)) and not isinstance(value, bool) else None
    for row in record["result"]["routes"]:
        scope = scope_by_key[row["route_key"]]
        role = labels[25] if scope.get("role") == "base" else labels[26]
        points = scope["points"]
        riders = (scope["before_input"].get("capacity_before" if scope.get("role") == "base" else "capacity_after")
                  if record["source_tool_key"] == "route_insert_advisor" else sum(p["passenger_count"] for p in points))
        routes.append([scope["plan_key"], row["route_id"], role,
            metric(row["before"].get("total_duration_s"), 60), metric(row["after"].get("total_duration_s"), 60),
            metric(row["delta"].get("total_duration_s"), 60), metric(row["before"].get("distance_m"), 1000),
            metric(row["after"].get("distance_m"), 1000), riders,
            row["after"].get("captured_at"), labels[23] if row["status"] == "verified" else labels[24]])
        evidence = dict(row.get("route_evidence") or {})
        valid = _verified(evidence, len(points))
        duration = distance = 0.0
        for index, point in enumerate(points):
            if valid and index:
                duration += evidence["legs"][index-1]["duration_s"]
                distance += evidence["legs"][index-1]["distance_m"]
            stops.append([scope["plan_key"], row["route_id"], role, index, point["address"], point["passenger_count"],
                labels[27] if point["is_depot"] else labels[28], metric(duration, 60) if valid else None,
                metric(distance, 1000) if valid else None])
    for sheet in wb:
        for cells in sheet:
            for cell in cells:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
        sheet.freeze_panes = "A2"
        sheet.sheet_view.showGridLines = False
        if sheet != summary:
            sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="136A75")
        for column in range(1, sheet.max_column + 1):
            sheet.column_dimensions[get_column_letter(column)].width = 24 if column == 1 else 20
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.print_title_rows = "1:1"
    stops.column_dimensions["E"].width = 48
    summary.column_dimensions["B"].width = 60
    for row in stops.iter_rows(min_row=2):
        row[4].alignment = Alignment(wrap_text=True, vertical="top")
    stream = BytesIO()
    wb.save(stream)
    return stream.getvalue()
