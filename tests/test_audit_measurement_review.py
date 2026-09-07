from copy import deepcopy
from dataclasses import asdict
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
import sys

import pytest
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "apps/backend"))
import audit_measurement_review as audit
import measurement_reviews as reviews
import planner_core as core
from runtime_store_sqlite import SqliteRuntimeStore
from test_measurement_review_api import fixture


def source(direction="To School"):
    config = core.PlannerConfig(service_direction=direction, minimum_vehicle_reduction=0,
        time_window_start="07:00" if direction == "To School" else "15:40",
        time_window_end="08:00" if direction == "To School" else "16:40", time_impact_limit_minutes=15)
    points = [dict(node_id=i, address=name, country="China", city="Shanghai", is_depot=i == 0,
                   passenger_count=count, lat=31.2 + i * .001, lng=121.4 + i * .001)
              for i, (name, count) in enumerate([("School", 0), ("A", 3), ("B", 2), ("Zero waypoint", 0)])]
    def route(route_id, nodes, frozen=False):
        if direction == "From School":
            nodes = list(reversed(nodes))
        return {"route_id": route_id, "vehicle_id": route_id, "nodes": nodes,
            "load": sum(points[i]["passenger_count"] for i in nodes), "bus_capacity": 42,
            "stop_service_time_s": sum(node != 0 for node in nodes[1:]) * 60,
            "time_s": (len(nodes) - 1) * 300 + sum(node != 0 for node in nodes[1:]) * 60,
            "distance_m": (len(nodes) - 1) * 500, "bus_type_name": "Bus",
            "leg_details": [{"duration_s": 300 + (60 if node else 0), "raw_osrm_duration_s": 300,
                             "distance_m": 500, "stop_service_s": 60 if node else 0} for node in nodes[1:]],
            "exception_role": "frozen_current" if frozen else "optimized",
            "final_route_traffic_gate": {"provider": "amap", "status": "passed", "target_duration_s": 3600,
                "verified_drive_duration_s": 300, "verified_total_duration_s": 360, "verified_distance_m": 500}}
    constraint = {"enabled": True, "mode": "hard", "strict_satisfied": True,
                  "expected_solver_stop_count": 3, "bounded_solver_stop_count": 3}
    def scenario(routes, is_current=False):
        return {"points": deepcopy(points), "routes": routes, "bus_count": len(routes), "stop_count": 3,
            "avg_route_duration_s": 9999, "avg_route_distance_m": 8888,
            "time_constraint": {} if is_current else deepcopy(constraint),
            "scenario_status": "passed", "traffic_gate": {"status": "passed", "provider": "amap"},
            "feasibility_report": {"status": "passed"}, "constraint_search_outcome": {"search_complete": True},
            "decision_metrics": {"evidence_complete": True, "affected_rider_count": 0}}
    structured = {"current_plan": scenario([route("R1", [1, 0]), route("R2", [2, 3, 0])], True),
        "time_constrained": scenario([route("Opt 1", [1, 2, 3, 0])]),
        "exception_preserving": scenario([route("R1", [1, 0], True), route("Opt 2", [2, 3, 0])]),
        "service_direction": direction, "planner_config": asdict(config), "input_address_count": 2}
    structured["exception_preserving"]["exception_preserving"] = {"enabled": True, "accepted": True, "frozen_route_ids": ["R1"], "attempts": [{"accepted": True}]}
    structured["current_plan_assessment"] = {"route_count": 2, "avg_route_duration_s": 9999,
        "total_duration_s": 19998, "total_distance_m": 17776, "overlong_route_count": 2,
        "recommendations": ["Original duration-based advice"],
        "route_summaries": [{"route_id": route["route_id"], "duration_s": 9999, "distance_m": 8888}
                            for route in structured["current_plan"]["routes"]]}
    return {"job_id": "source", "owner_email": "owner@example.test", "status": "succeeded",
        "metadata": {"job_kind": "route_audit"}, "config": asdict(config),
        "result": {"structured_results": structured, "planner_config": asdict(config)}}


class Provider:
    calls = []
    direct_a = 600
    def __init__(self, name, *, departure_time, api_call_limit):
        assert name == "amap" and departure_time is None
        self.provider, self.planner = name, SimpleNamespace(AMAP_KEY="unit-test")
        self.state = {"api_calls": 0, "api_call_limit": api_call_limit, "cache_hits": 0}
    def route(self, points, **kwargs):
        pairs = [(a["node_id"], b["node_id"]) for a, b in zip(points, points[1:])]
        if self.state["api_calls"] + len(pairs) > self.state["api_call_limit"]:
            raise RuntimeError("Budget exhausted")
        self.state["api_calls"] += len(pairs)
        self.calls.append(pairs)
        durations = [self.direct_a if set(pair) == {0, 1} else 1800 if set(pair) == {1, 2} else 300 for pair in pairs]
        geometry = [[point["lng"], point["lat"]] for point in points]
        return {"evidence_version": reviews.EVIDENCE_VERSION, "provider": "amap", "source": "amap_adjacent_legs",
            "status": "verified", "complete": True, "called_at": datetime.now(timezone.utc).isoformat(),
            "duration_s": sum(durations), "distance_m": len(pairs) * 500, "leg_durations_s": durations,
            "geometry": geometry, "legs": [{"duration_s": duration, "distance_m": 500, "geometry": geometry[i:i + 2]}
                for i, duration in enumerate(durations)]}


@pytest.fixture(autouse=True)
def isolated(monkeypatch):
    Provider.calls, Provider.direct_a = [], 600
    monkeypatch.setattr(core, "FINAL_ROUTE_TRAFFIC_VERIFICATION_ENABLED", True)
    monkeypatch.setattr(core, "load_legacy_planner", lambda: pytest.fail("No live provider initialization"))
    monkeypatch.setattr(core, "_final_route_stats", lambda *args, **kwargs: pytest.fail("No unbounded measurement path"))


def create(store, original=None, limit=100):
    original = original or source()
    store.upsert_job(original)
    request = audit.build_audit_review_request(original, requested_by="admin@example.test", request_key="audit-full", provider_call_limit=limit)
    return store.create_route_measurement_review(request, queue_scope="test")


@pytest.mark.parametrize("direction", ["To School", "From School"])
def test_full_audit_revalidates_time_impact_without_changing_solver_inputs(tmp_path, direction):
    store = SqliteRuntimeStore(tmp_path / "runtime.sqlite")
    original = source(direction)
    row = create(store, original)
    before = deepcopy(store.get_job("source"))
    completed = reviews.execute_saved_review(store, row["review_id"], "worker", provider_factory=Provider)
    assert completed["status"] == "succeeded", completed
    result = completed["result"]
    assert result["time_window_revalidated"] and result["time_impact_revalidated"]
    assert not result["solver_rerun"] and not result["minimum_fleet_revalidated"]
    structured = result["audit_result"]["structured_results"]
    strict = structured["time_constrained"]
    assert strict["traffic_gate"]["status"] == "passed"
    expected_limit = core.effective_route_duration_limit_minutes(core.build_planner_config(original["config"]))
    assert strict["traffic_gate"]["solver_target_duration_minutes"] == expected_limit
    assert strict["traffic_gate"]["target_duration_minutes"] == (expected_limit if direction == "To School" else 60)
    assert strict["final_time_impact_gate"]["status"] == "failed"
    assert strict["final_time_impact_gate"]["max_adverse_minutes"] == pytest.approx(32)
    assert strict["final_time_impact_gate"]["over_limit_rider_count"] == 3
    assert strict["scenario_status"] == "rejected" and strict["feasibility_report"]["status"] == "failed"
    assert strict["constraint_search_outcome"]["search_complete"] is False
    assert "decision_metrics" not in strict
    assert strict["avg_route_duration_s"] == (2520 if direction == "To School" else 2580)
    assert structured["exception_preserving"]["exception_feasible"] is True
    assert structured["time_constrained_optimization"] == strict
    assessment = result["audit_result"]["current_plan_assessment"]
    assert assessment["total_duration_s"] == (1260 if direction == "To School" else 1380)
    assert assessment["total_distance_m"] == 1500 and assessment["overlong_route_count"] == 0
    assert "recommendations" not in assessment
    assert store.get_job("source") == before
    audit.validate_audit_result(row["request"], result, terminal=True)


def test_protected_plan_keeps_frozen_exception_but_reports_all_route_failure(tmp_path):
    Provider.direct_a = 4200
    store = SqliteRuntimeStore(tmp_path / "runtime.sqlite")
    row = create(store)
    completed = reviews.execute_saved_review(store, row["review_id"], "worker", provider_factory=Provider)
    assert completed["status"] == "succeeded", completed
    protected = completed["result"]["audit_result"]["structured_results"]["exception_preserving"]
    assert protected["traffic_gate"]["all_routes_status"] == "failed"
    assert protected["traffic_gate"]["status"] == "passed"
    assert protected["exception_preserving"]["candidate_failure_summary"]["failed_route_ids"] == ["R1"]
    assert protected["exception_feasible"] and protected["exception_preserving"]["frozen_route_ids"] == ["R1"]


def test_exhausted_budget_never_retains_old_pass_or_map_evidence(tmp_path):
    store = SqliteRuntimeStore(tmp_path / "runtime.sqlite")
    row = create(store, limit=1)
    completed = reviews.execute_saved_review(store, row["review_id"], "worker", provider_factory=Provider)
    assert completed["status"] == "needs_review", completed
    result = completed["result"]
    assert result["provider_api_calls"] == completed["api_calls"] == 1
    assert not result["time_window_revalidated"] and not result["time_impact_revalidated"]
    strict = result["audit_result"]["structured_results"]["time_constrained"]
    assert strict["scenario_status"] != "passed"
    assert strict["final_time_impact_gate"]["status"] == "unavailable"
    assert strict["avg_route_duration_s"] is None
    assert strict["routes"][0]["route_evidence"]["duration_s"] is None
    assert strict["routes"][0]["final_route_traffic_gate"]["status"] == "unavailable"
    assessment = result["audit_result"]["current_plan_assessment"]
    assert assessment["total_duration_s"] is None and assessment["overlong_route_count"] is None


def test_pause_resume_reuses_review_snapshots_and_keeps_budget(tmp_path, monkeypatch):
    store = SqliteRuntimeStore(tmp_path / "runtime.sqlite")
    row = create(store)
    save = store.save_review_measurement_snapshot
    interrupted = []
    def pause(*args):
        saved = save(*args)
        if not interrupted:
            interrupted.append(True)
            store.pause_route_measurement_review(row["review_id"], yielding=True)
        return saved
    monkeypatch.setattr(store, "save_review_measurement_snapshot", pause)
    first = reviews.execute_saved_review(store, row["review_id"], "first", provider_factory=Provider)
    assert first["status"] == "yielding" and first["api_calls"] == 1
    store.finish_route_measurement_worker(row["review_id"], "first")
    second = reviews.execute_saved_review(store, row["review_id"], "second", provider_factory=Provider)
    assert second["status"] == "succeeded", second
    assert Provider.calls.count([(1, 0)]) == 1
    assert second["api_calls"] == sum(len(pairs) for pairs in Provider.calls)


@pytest.mark.parametrize("change", ["missing_config", "fractional_config", "missing_current", "invalid_order", "dwell", "capacity"])
def test_ambiguous_original_inputs_cannot_be_silently_replaced(change):
    original = source()
    config = original["result"]["planner_config"]
    current = original["result"]["structured_results"]["current_plan"]
    if change == "missing_config": config.pop("time_impact_limit_minutes")
    if change == "fractional_config": config["stop_service_minutes"] = 1.5
    if change == "missing_current": current["routes"] = []
    if change == "invalid_order": current["routes"][0]["nodes"] = [0, 1]
    if change == "dwell": current["routes"][0]["stop_service_time_s"] += 60
    if change == "capacity": current["routes"][0].pop("bus_capacity")
    with pytest.raises(ValueError):
        audit.build_audit_review_request(original, requested_by="admin@example.test", request_key="x", provider_call_limit=100)
    assert not Provider.calls


def test_empty_saved_candidate_remains_unresolved_not_infeasibility_proof(tmp_path):
    original = source()
    original["result"]["structured_results"]["time_constrained"] = {"points": [], "routes": [], "scenario_status": "infeasible"}
    store = SqliteRuntimeStore(tmp_path / "runtime.sqlite")
    row = create(store, original)
    completed = reviews.execute_saved_review(store, row["review_id"], "worker", provider_factory=Provider)
    assert completed["status"] == "needs_review", completed
    assert completed["result"]["audit_result"]["structured_results"]["time_constrained"]["scenario_status"] == "unresolved"


def test_legacy_strict_route_without_id_uses_native_identity_without_mutation(tmp_path):
    original = source()
    original["result"]["structured_results"]["time_constrained"]["routes"][0].pop("route_id")
    store = SqliteRuntimeStore(tmp_path / "runtime.sqlite")
    row = create(store, original)
    assert "time_constrained:Bus 1" in [scope["route_key"] for scope in row["request"]["routes"]]
    completed = reviews.execute_saved_review(store, row["review_id"], "worker", provider_factory=Provider)
    assert completed["status"] == "succeeded", completed
    strict = completed["result"]["audit_result"]["structured_results"]["time_constrained"]
    assert "route_id" not in strict["routes"][0]
    assert strict["routes"][0]["final_route_traffic_gate"]["route_id"] == "Bus 1"
    assert strict["final_time_impact_gate"]["max_adverse_minutes"] == pytest.approx(32)
    assert store.get_job("source")["result"] == original["result"]
    audit.validate_audit_result(row["request"], completed["result"], terminal=True)


def test_native_identity_collision_is_not_silently_matched():
    original = source()
    routes = original["result"]["structured_results"]["current_plan"]["routes"]
    routes[0].pop("route_id")
    routes[1]["route_id"] = "Bus 1"
    with pytest.raises(ValueError, match="ambiguous"):
        audit.build_audit_review_request(original, requested_by="admin@example.test", request_key="x", provider_call_limit=100)


def test_full_audit_terminal_integrity_rejects_changed_vehicles_and_points(tmp_path):
    store = SqliteRuntimeStore(tmp_path / "runtime.sqlite")
    row = create(store)
    completed = reviews.execute_saved_review(store, row["review_id"], "worker", provider_factory=Provider)
    changed = deepcopy(completed["result"])
    changed["audit_result"]["structured_results"]["time_constrained"]["routes"][0]["nodes"] = [3, 2, 1, 0]
    changed["audit_result"]["structured_results"]["time_constrained_optimization"] = deepcopy(changed["audit_result"]["structured_results"]["time_constrained"])
    changed["audit_result"]["time_constrained_optimization"] = deepcopy(changed["audit_result"]["structured_results"]["time_constrained"])
    with pytest.raises(ValueError, match="cannot change routes"):
        audit.validate_audit_result(row["request"], changed, terminal=True)
    changed = deepcopy(completed["result"])
    changed["audit_result"]["structured_results"]["current_plan"]["points"][1]["lat"] += .001
    changed["audit_result"]["structured_results"]["current_plan_scenario"] = deepcopy(changed["audit_result"]["structured_results"]["current_plan"])
    changed["audit_result"]["current_plan_scenario"] = deepcopy(changed["audit_result"]["structured_results"]["current_plan"])
    with pytest.raises(ValueError, match="cannot replace"):
        audit.validate_audit_result(row["request"], changed, terminal=True)


def test_disabled_validation_does_not_certify_a_correction(tmp_path, monkeypatch):
    monkeypatch.setattr(core, "FINAL_ROUTE_TRAFFIC_VERIFICATION_ENABLED", False)
    store = SqliteRuntimeStore(tmp_path / "runtime.sqlite")
    row = create(store)
    completed = reviews.execute_saved_review(store, row["review_id"], "worker", provider_factory=Provider)
    assert completed["status"] == "needs_review", completed
    assert not Provider.calls
    for key in audit.SCENARIOS:
        scenario = completed["result"]["audit_result"]["structured_results"][key]
        assert scenario["scenario_status"] != "passed"
        assert scenario["traffic_gate"]["status"] == "unavailable"


@pytest.mark.parametrize("change", ["missing_scenario", "frozen_identity"])
def test_saved_scenarios_and_frozen_identity_cannot_be_inferred(change):
    original = source()
    structured = original["result"]["structured_results"]
    if change == "missing_scenario":
        structured.pop("time_constrained")
    else:
        structured["exception_preserving"]["exception_preserving"]["frozen_route_ids"] = ["wrong"]
    with pytest.raises(ValueError):
        audit.build_audit_review_request(original, requested_by="admin@example.test", request_key="x", provider_call_limit=100)


def test_capacity_and_vehicle_reduction_still_gate_corrected_candidates(tmp_path):
    original = source()
    original["result"]["planner_config"]["minimum_vehicle_reduction"] = 1
    original["result"]["structured_results"]["time_constrained"]["routes"][0]["load"] = 50
    store = SqliteRuntimeStore(tmp_path / "runtime.sqlite")
    row = create(store, original)
    completed = reviews.execute_saved_review(store, row["review_id"], "worker", provider_factory=Provider)
    assert completed["status"] == "succeeded", completed
    structured = completed["result"]["audit_result"]["structured_results"]
    assert "physical_capacity" in structured["time_constrained"]["feasibility_report"]["failure_reasons"]
    assert "vehicle_savings_target" in structured["exception_preserving"]["feasibility_report"]["failure_reasons"]


def test_audit_mode_is_explicit_and_source_bound(fixture):
    client, store, _schedule = fixture
    store.upsert_job(source())
    body = {"mode": "full_audit", "request_key": "audit-api", "provider_call_limit": 100, "confirm_provider_calls": True}
    url = "/api/jobs/source/measurement-reviews"
    assert client.post(url, json={**body, "route_keys": ["current_plan:R1"]}).status_code == 422
    assert client.post(url, json=body, headers={"X-BRP-User-Email": "owner@example.test"}).status_code == 403
    response = client.post(url, json=body)
    assert response.status_code == 200, response.text
    assert response.json()["request"]["mode"] == "full_audit"


def test_corrected_map_and_native_workbook_use_saved_measurements(fixture, monkeypatch):
    client, store, _schedule = fixture
    original = source()
    store.upsert_job(original)
    row = create(store, original)
    base = f"/api/jobs/source/measurement-reviews/{row['review_id']}"
    assert client.get(base + "/map-data/time_constrained").status_code == 409
    completed = reviews.execute_saved_review(store, row["review_id"], "worker", provider_factory=Provider)
    assert completed["status"] == "succeeded", completed
    before_calls = deepcopy(Provider.calls)
    import backend_service
    monkeypatch.setattr(backend_service, "_amap_display_geometry_for_route", lambda *args, **kwargs: pytest.fail("No display-route fetching"))
    response = client.get(base + "/map-data/time_constrained")
    assert response.status_code == 200, response.text
    data = response.json()
    assert data["routes"][0]["verified_total_duration_s"] == 2520
    assert data["routes"][0]["duration_s"] == 2520
    assert data["routes"][0]["distance_m"] == 1500
    assert data["summary"]["time_impact"]["over_acceptance_rider_count"] == 3
    assert data["summary"]["time_impact"]["max_adverse_delta_minutes"] == pytest.approx(32)
    denied = client.get(base + "/map-data/time_constrained", headers={"X-BRP-User-Email": "stranger@example.test"})
    assert denied.status_code == 403
    report = client.get(base + "/export")
    assert report.status_code == 200, report.text
    book = load_workbook(BytesIO(report.content))
    assert {"Acceptance Review", "Road Comparison", "Strict Summary", "Protected Summary"}.issubset(book.sheetnames)
    assert book["Road Comparison"].column_dimensions["I"].width == 44
    assert list(book["Acceptance Review"].values)[5][:6] == ("Strict Plan", "Passed", "Not passed", "Passed", "All routes", "Not passed")
    assert list(book["Acceptance Review"].values)[6][4] == "Non-frozen routes"
    strict_comparison = next(row for row in book["Road Comparison"].iter_rows(min_row=5, values_only=True) if row[0] == "Strict Plan")
    assert strict_comparison[1] == data["routes"][0]["id"]
    native_bytes, error = backend_service._build_time_impact_workbook_export(audit.audit_review_record(completed), "time_constrained")
    assert not error
    native_book = load_workbook(BytesIO(native_bytes))
    for sheet in native_book:
        assert list(book[f"Strict {sheet.title}"].values) == list(sheet.values)
    assert Provider.calls == before_calls


def test_incomplete_correction_map_cannot_show_old_planning_totals(fixture, monkeypatch):
    client, store, _schedule = fixture
    row = create(store, limit=1)
    completed = reviews.execute_saved_review(store, row["review_id"], "worker", provider_factory=Provider)
    assert completed["status"] == "needs_review"
    import backend_service
    monkeypatch.setattr(backend_service, "_build_job_map_payload", lambda *args, **kwargs: pytest.fail("Incomplete evidence must not enter legacy map fallback"))
    base = f"/api/jobs/source/measurement-reviews/{row['review_id']}"
    for key in audit.SCENARIOS:
        assert client.get(base + f"/map-data/{key}").status_code == 409
    report = client.get(base + "/export")
    assert report.status_code == 200
    book = load_workbook(BytesIO(report.content))
    assert "Strict Availability" in book and "Strict Routes" not in book
    assert "Protected Availability" in book and "Protected Routes" not in book
    assert book["Strict Availability"]["A5"].value == "Needs review"
    assert book["Road Comparison"]["D7"].value is None
