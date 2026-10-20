from copy import deepcopy
from datetime import timedelta
from io import BytesIO
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from test_paired_return import NOW, source
from test_audit_paired_return import audit_api
from audit_return import bind_morning_scenarios, related_pair_records, run_audit_return, scenario_payload
from paired_return import build_return_workbook, ensure_return_due
import backend_service as service
import backend_job_runner as runner
import api_app


def pair_fixture():
    prepared = source()
    points = [{**p, "node_id": index} for index, p in enumerate(prepared["original_points"])]
    scenarios = {}
    for key, routes, status in [
        ("current_plan", [[1, 2, 0]], "passed"),
        ("time_constrained", [[1, 0], [2, 0]], "rejected"),
        ("exception_preserving", [[2, 1, 0]], "passed"),
    ]:
        scenarios[key] = {"points": deepcopy(points), "scenario_status": status, "routes": [
            {"vehicle_id": i + 1, "nodes": nodes, "load": sum(points[n]["passenger_count"] for n in nodes)}
            for i, nodes in enumerate(routes)]}
    ids = {"am": "test-am", "pm": "test-pm"}
    common = {"pair_id": "test", "paired_job_ids": ids, "job_kind": "route_audit", "job_queue_scope": "staging"}
    am = {"job_id": ids["am"], "owner_email": "owner@example.com", "status": "succeeded",
          "created_at": NOW.isoformat(), "finished_at": NOW.isoformat(),
          "prepared_payload": prepared, "result": {"structured_results": scenarios},
          "metadata": {**common, "period": "am"}}
    pm = {"job_id": ids["pm"], "owner_email": am["owner_email"], "status": "scheduled",
          "created_at": NOW.isoformat(), "scheduled_start_at": (NOW - timedelta(minutes=10)).isoformat(),
          "metadata": {**common, "period": "pm", "measurement_kind": "pm_return",
                       "scheduled_start_at": (NOW - timedelta(minutes=10)).isoformat(),
                       "schedule_expires_at": (NOW - timedelta(minutes=1)).isoformat(),
                       "analysis_config": {"stop_service_minutes": 1, "final_time_validation_mode": "legacy"}},
          "prepared_payload": {}}
    return am, pm


def bind(am, pm, now=NOW):
    return bind_morning_scenarios(pm, {am["job_id"]: am, pm["job_id"]: pm}.get, now)


def test_wait_then_bind_all_three_including_rejected_not_original():
    am, pm = pair_fixture()
    am["status"] = "running"
    assert bind(am, pm) is False
    assert pm["prepared_payload"] == {}
    am["status"] = "succeeded"
    assert bind(am, pm)
    snapshots = pm["prepared_payload"]["return_scenarios"]
    assert list(snapshots) == ["current_plan", "time_constrained", "exception_preserving"]
    assert snapshots["time_constrained"]["source_status"] == "rejected"
    assert [s["address"] for s in snapshots["exception_preserving"]["prepared_payload"]["current_plan"]["stops"]] == ["School", "A", "B"]
    assert pm["metadata"]["return_delayed"] is True
    ensure_return_due(pm["metadata"], NOW)
    frozen = deepcopy(pm["prepared_payload"])
    am["result"]["structured_results"]["exception_preserving"]["routes"].clear()
    assert bind(am, pm)
    assert pm["prepared_payload"] == frozen


@pytest.mark.parametrize("status", ["failed", "canceled"])
def test_failed_parent_has_no_original_fallback(status):
    am, pm = pair_fixture()
    am["status"] = status
    with pytest.raises(ValueError, match=status):
        bind(am, pm)
    assert pm["prepared_payload"] == {}


def test_missing_parent_and_expired_post_completion_window():
    am, pm = pair_fixture()
    with pytest.raises(ValueError, match="missing"):
        bind_morning_scenarios(pm, lambda _: None, NOW)
    with pytest.raises(ValueError, match="expired"):
        bind(am, pm, NOW + timedelta(minutes=16))


def test_invalid_riders_missing_scenario_never_substituted():
    am, pm = pair_fixture()
    am["result"]["structured_results"]["time_constrained"]["routes"][0]["load"] = 99
    del am["result"]["structured_results"]["exception_preserving"]
    assert bind(am, pm)
    scenarios = pm["prepared_payload"]["return_scenarios"]
    assert "rider assignments" in scenarios["time_constrained"]["error"]
    assert "no route solution" in scenarios["exception_preserving"]["error"]
    assert scenarios["current_plan"].get("prepared_payload")


def test_three_scenario_measurements_export_and_independent_highlights(monkeypatch):
    import direct_school_analysis as analysis
    am, pm = pair_fixture()
    bind(am, pm)
    calls = []

    class Provider:
        def __init__(self, *args, **kwargs):
            self.state = {"api_calls": 0}

        def route(self, points):
            calls.append([p["address"] for p in points])
            self.state["api_calls"] += 1
            seconds = 3500 if points[-1]["address"] == "B" else 1000
            return {"complete": True, "duration_s": seconds, "distance_m": 1234,
                    "leg_durations_s": [seconds / (len(points) - 1)] * (len(points) - 1),
                    "geometry": [[121, 31], [121.1, 31.1]]}

    monkeypatch.setattr(analysis, "_osrm_leg", lambda *args: {"duration_s": 100, "distance_m": 100})
    checkpoints = []
    result = run_audit_return(pm, provider_factory=Provider, now=lambda: NOW, checkpoint=checkpoints.append)
    assert result["status"] == "complete"
    assert calls == [["School", "B", "A"], ["School", "A"], ["School", "B"], ["School", "A", "B"]]
    assert result["progress"]["provider_api_calls"] == 4
    assert result["scenarios"]["exception_preserving"]["summary"]["under_one_hour_count"] == 0
    assert result["scenarios"]["time_constrained"]["summary"]["rider_count"] == 2
    assert result["scenarios"]["time_constrained"]["source_status"] == "rejected"
    book = load_workbook(BytesIO(build_return_workbook({**pm, "result": result})))
    assert len(book.sheetnames) == 7
    assert book["Strict Return routes"].max_row == 7
    assert book["Protected Return routes"]["G6"].value == "No / 否"
    assert book["Scenarios"]["B6"].value == "rejected"
    assert checkpoints[-1]["progress"]["completed"] == 4


@pytest.mark.parametrize("field", ["owner", "scope", "reciprocal"])
def test_pair_links_cannot_delete_unrelated_job(field):
    am, pm = pair_fixture()
    if field == "owner":
        pm["owner_email"] = "other@example.com"
    elif field == "scope":
        pm["metadata"]["job_queue_scope"] = "prod"
    else:
        pm["metadata"]["paired_job_ids"] = {"am": "unrelated", "pm": pm["job_id"]}
    with pytest.raises(ValueError, match="mismatch"):
        related_pair_records(am, {am["job_id"]: am, pm["job_id"]: pm}.get)


@pytest.mark.parametrize("period", ["am", "pm"])
def test_api_delete_either_member_removes_both_and_retains_unrelated(audit_api, monkeypatch, period):
    client, store, payload, context, _ = audit_api
    pair = client.post("/api/workbooks/submit", json=payload).json()["paired_jobs"]
    am, pm = [store.get_job(item["job_id"]) for item in pair]
    unrelated = {**deepcopy(am), "job_id": "unrelated", "metadata": {}}
    store.upsert_job(unrelated)
    monkeypatch.setattr(service.JOB_QUEUE, "schedule_queued_jobs", lambda: None)
    target = am if period == "am" else pm
    response = client.delete(f"/api/jobs/{target['job_id']}")
    assert response.status_code == 200, response.text
    assert set(response.json()["deleted_job_ids"]) == {am["job_id"], pm["job_id"]}
    assert store.get_job(am["job_id"]) is None
    assert store.get_job(pm["job_id"]) is None
    assert store.get_job("unrelated")
    assert service.JOB_STORE.release_due_scheduled_jobs() == []


def test_scheduler_waits_then_releases_without_original_snapshot(audit_api, monkeypatch):
    _, store, _, _, _ = audit_api
    am, pm = pair_fixture()
    # Use the scheduler's own clock, not the fixed measurement fixture date.
    from datetime import datetime, timezone
    current = datetime.now(timezone.utc)
    pm["scheduled_start_at"] = (current - timedelta(minutes=10)).isoformat()
    pm["metadata"]["scheduled_start_at"] = pm["scheduled_start_at"]
    am["finished_at"] = current.isoformat()
    am["status"] = "running"
    am["metadata"]["job_queue_scope"] = service.JOB_QUEUE_SCOPE
    pm["metadata"]["job_queue_scope"] = service.JOB_QUEUE_SCOPE
    for record in (am, pm):
        store.upsert_job(record)
    assert service.JOB_STORE.release_due_scheduled_jobs() == []
    assert store.get_job(pm["job_id"])["metadata"]["return_dependency_status"] == "waiting_for_morning"
    am["status"] = "succeeded"
    store.upsert_job(am)
    released = service.JOB_STORE.release_due_scheduled_jobs()
    assert len(released) == 1
    assert set(released[0]["prepared_payload"]["return_scenarios"]) == {"current_plan", "time_constrained", "exception_preserving"}
    assert service.JOB_STORE.release_due_scheduled_jobs() == []


def test_group_delete_transaction_rolls_back_on_failure(tmp_path, monkeypatch):
    from runtime_store_sqlite import SqliteRuntimeStore
    store = SqliteRuntimeStore(tmp_path / "atomic.sqlite")
    am, pm = pair_fixture()
    for record in (am, pm):
        store.upsert_job(record)
    original = store._delete_history_group_items

    def fail_second(conn, scopes, job_id):
        if job_id == pm["job_id"]:
            raise RuntimeError("disk failure")
        return original(conn, scopes, job_id)

    monkeypatch.setattr(store, "_delete_history_group_items", fail_second)
    with pytest.raises(RuntimeError):
        store.delete_jobs([am["job_id"], pm["job_id"]])
    assert store.get_job(am["job_id"])
    assert store.get_job(pm["job_id"])


def test_api_group_delete_stops_both_workers_and_background_verification(audit_api, monkeypatch):
    client, store, payload, _, _ = audit_api
    pair = client.post("/api/workbooks/submit", json=payload).json()["paired_jobs"]
    am, pm = [store.get_job(item["job_id"]) for item in pair]
    for index, record in enumerate((am, pm)):
        record.update(status="running", worker_pid=900001 + index)
        store.upsert_job(record)
    verification = {"verification_id": "test-deep", "parent_job_id": am["job_id"],
                    "scenario_key": "exception_preserving", "status": "running", "worker_pid": 900003}
    store.upsert_deep_verification(verification)
    stopped = []
    import job_queue
    monkeypatch.setattr(job_queue, "terminate_worker_process", stopped.append)
    monkeypatch.setattr(service.JOB_QUEUE, "schedule_queued_jobs", lambda: None)
    response = client.delete(f"/api/jobs/{am['job_id']}")
    assert response.status_code == 200
    assert set(stopped) == {900001, 900002, 900003}
    assert store.get_deep_verification("test-deep") is None
    assert all(store.get_job(record["job_id"]) is None for record in (am, pm))


def test_api_pair_mismatch_does_not_delete_either_member(audit_api, monkeypatch):
    client, store, payload, _, _ = audit_api
    pair = client.post("/api/workbooks/submit", json=payload).json()["paired_jobs"]
    am, pm = [store.get_job(item["job_id"]) for item in pair]
    pm["owner_email"] = "other@example.com"
    store.upsert_job(pm)
    response = client.delete(f"/api/jobs/{am['job_id']}")
    assert response.status_code == 409
    assert store.get_job(am["job_id"])
    assert store.get_job(pm["job_id"])


def test_worker_defers_until_parent_without_calling_provider(monkeypatch):
    import sys
    am, pm = pair_fixture()
    am["status"] = "running"
    state = {am["job_id"]: am, pm["job_id"]: pm}
    monkeypatch.setattr(runner, "_load_job", lambda key: deepcopy(state[key]))
    monkeypatch.setattr(runner, "_save_job", lambda record: state.update({record["job_id"]: deepcopy(record)}))
    monkeypatch.setattr(runner, "_runtime_store", lambda: SimpleNamespace(get_job=state.get))
    monkeypatch.setattr(runner, "_release_concurrency_slot", lambda: None)
    monkeypatch.setattr(sys, "argv", ["backend_job_runner.py", pm["job_id"]])
    assert runner.main() == 0
    assert state[pm["job_id"]]["status"] == "scheduled"
    assert state[pm["job_id"]]["started_at"] is None
