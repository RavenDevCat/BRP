from copy import deepcopy
from datetime import datetime, timedelta, timezone
from io import BytesIO
from pathlib import Path
import sys

from openpyxl import load_workbook
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "apps/backend"))
import direct_school_analysis as analysis
import full_measurement_review as full
import measurement_reviews as reviews
from runtime_store_sqlite import SqliteRuntimeStore
from test_measurement_review_api import fixture


def source(direction="To School"):
    counts = {"School": 0, "A": 4, "B": 3, "C": 2, "D": 1, "Zero waypoint": 0}
    points = [dict(country="China", city="Shanghai", address=name, lat=31.2 + i * .001,
                   lng=121.4 + i * .001, passenger_count=count) for i, (name, count) in enumerate(counts.items())]
    names = ["A", "B", "C", "D", "Zero waypoint", "School"]
    if direction == "From School":
        names.reverse()
    stops = [dict(country="China", city="Shanghai", address=name, passenger_count=counts[name],
                  route_id="R1", stop_sequence=i, is_depot=name == "School") for i, name in enumerate(names)]
    config = {**analysis.DEFAULT_ANALYSIS_CONFIG, "service_direction": direction, "far_duration_minutes": 60,
              "time_window_start": "07:15" if direction == "To School" else "15:40",
              "time_window_end": "08:00" if direction == "To School" else "16:25"}
    old_stops = [{"stop_key": analysis._stop_key(stop), "address": stop["address"], "riders": stop["passenger_count"],
                  "operational_category": "route_only_over_limit"} for stop in stops if not stop["is_depot"]]
    return {"job_id": "source", "owner_email": "owner@example.test", "status": "succeeded",
        "metadata": {"job_kind": "direct_school_analysis"},
        "prepared_payload": {"input_records": points, "original_points": deepcopy(points),
                             "current_plan": {"service_direction": direction, "stops": stops}},
        "result": {"parameters": config, "provider": "amap", "service_direction": direction, "stops": old_stops,
                   "routes": [{"route_id": "R1", "provider_duration_min": 120, "total_duration_min": 125,
                               "provider_distance_km": 90, "status": "resolved"}],
                   "summary": {"direct_over_limit_rider_count": 0, "route_only_over_limit_rider_count": 10,
                               "additional_removal_rider_count": 0, "routes_over_window_final_count": 1}}}


def osrm(origin, destination, _cache):
    return {"duration_s": 300, "distance_m": 2000,
            "geometry": [[origin["lng"], origin["lat"]], [destination["lng"], destination["lat"]]]}


class Provider:
    calls = []

    def __init__(self, provider, *, departure_time, api_call_limit):
        assert provider == "amap" and departure_time is None
        self.provider = provider
        self.state = {"api_calls": 0, "api_call_limit": api_call_limit, "cache_hits": 0}

    def route(self, points, **_kwargs):
        names = [point["address"] for point in points]
        count = len(points) - 1
        if self.state["api_calls"] + count > self.state["api_call_limit"]:
            raise RuntimeError("Budget exhausted")
        self.state["api_calls"] += count
        self.calls.append(names)
        durations = [1000] * count
        if count == 1 and "School" in names:
            name = names[0] if names[1] == "School" else names[1]
            durations = [{"A": 4200, "B": 2000, "C": 1500, "D": 1200, "Zero waypoint": 1000}[name]]
        geometry = [[point["lng"], point["lat"]] for point in points]
        result = {"evidence_version": full.EVIDENCE_VERSION, "provider": "amap", "complete": True,
            "status": "verified", "duration_s": sum(durations), "distance_m": count * 2000,
            "leg_durations_s": durations, "leg_distances_m": [2000] * count,
            "called_at": datetime.now(timezone.utc).isoformat(), "geometry": geometry,
            "legs": [{"duration_s": value, "distance_m": 2000, "geometry": geometry[i:i + 2]}
                     for i, value in enumerate(durations)]}
        self.state["last_route_evidence"] = deepcopy(result)
        return result


@pytest.fixture(autouse=True)
def isolate(monkeypatch):
    Provider.calls = []
    monkeypatch.setattr(analysis, "_osrm_leg", osrm)
    monkeypatch.setattr(analysis.planner_core, "load_legacy_planner", lambda: pytest.fail("No real provider initialization"))


def create(store, original=None, *, limit=100, key="full-test"):
    original = original or source()
    store.upsert_job(original)
    payload = full.build_full_review_request(original, requested_by="admin@example.test", request_key=key, provider_call_limit=limit)
    return store.create_route_measurement_review(payload, queue_scope="test")


@pytest.mark.parametrize("direction", ["To School", "From School"])
def test_full_pipeline_recomputes_all_three_groups_and_keeps_original(tmp_path, direction):
    store = SqliteRuntimeStore(tmp_path / "runtime.sqlite")
    original = source(direction)
    row = create(store, original)
    before = deepcopy(store.get_job("source"))
    result = reviews.execute_saved_review(store, row["review_id"], "worker", provider_factory=Provider)
    assert result["status"] == "succeeded", result
    report = result["result"]
    corrected = report["analysis_result"]
    assert report["student_classification_recomputed"] and report["classification_complete"]
    assert report["time_window_revalidated"] and report["scope"] == "full_direct_school_result"
    conclusion = corrected["operational_conclusion"]
    assert conclusion["direct_over_limit"]["rider_count"] == 4
    assert conclusion["route_only_over_limit"]["rider_count"] == 3
    assert conclusion["additional_removal"]["rider_count"] == 2
    recovered = corrected["route_window_analysis"][0]
    assert recovered["original_riders"] == 10 and recovered["final_riders"] == 1
    assert recovered["additional_removals"][0]["address"] == "C"
    assert recovered["additional_removal_strategy"] == analysis.ADDITIONAL_REMOVAL_STRATEGY
    assert corrected["routes"][0]["stop_count"] == 5
    assert len(corrected["stops"]) == 5
    expected_order = [stop["address"] for stop in original["prepared_payload"]["current_plan"]["stops"]]
    assert expected_order in Provider.calls
    b = next(stop for stop in corrected["stops"] if stop["address"] == "B")
    assert b["route_contexts"][0]["estimated_current_ride_min"] == pytest.approx(70.67)
    assert corrected["parameters"] == {**original["result"]["parameters"], "provider_call_limit": 100}
    assert report["conclusion_changes"]["direct_over_limit_rider_count"] == {"before": 0, "after": 4, "delta": 4}
    assert store.get_job("source") == before


def test_full_review_matches_existing_analysis_not_a_second_algorithm(tmp_path):
    original = source()
    config = {**original["result"]["parameters"], "provider_call_limit": 100}
    expected = analysis.run_direct_school_analysis(original["prepared_payload"], config, run_seed="full-test", provider_factory=Provider)
    store = SqliteRuntimeStore(tmp_path / "runtime.sqlite")
    row = create(store, original)
    actual = reviews.execute_saved_review(store, row["review_id"], "worker", provider_factory=Provider)["result"]["analysis_result"]
    assert actual["operational_conclusion"] == expected["operational_conclusion"]
    assert [(r["stop_key"], r["operational_category"], r.get("route_contexts")) for r in actual["stops"]] == [
        (r["stop_key"], r["operational_category"], r.get("route_contexts")) for r in expected["stops"]]


def test_budget_exhaustion_keeps_unknowns_in_complete_scope(tmp_path):
    store = SqliteRuntimeStore(tmp_path / "runtime.sqlite")
    row = create(store, limit=1)
    result = reviews.execute_saved_review(store, row["review_id"], "worker", provider_factory=Provider)
    assert result["status"] == "needs_review", result
    report = result["result"]
    assert result["api_calls"] == 1 and report["provider_api_calls"] == 1
    assert not report["classification_complete"] and not report["time_window_revalidated"]
    assert len(report["analysis_result"]["stops"]) == 5
    assert report["analysis_result"]["route_window_analysis"][0]["status"] == "data_review"
    assert report["routes"][0]["after"]["total_duration_s"] is None


def test_resume_reuses_owned_fresh_measurements_without_resetting_budget(tmp_path, monkeypatch):
    store = SqliteRuntimeStore(tmp_path / "runtime.sqlite")
    row = create(store)
    rid = row["review_id"]
    save = store.save_review_measurement_snapshot
    paused = []
    def pause_after_snapshot(*args):
        value = save(*args)
        if not paused:
            paused.append(True)
            store.pause_route_measurement_review(rid, yielding=True)
        return value
    monkeypatch.setattr(store, "save_review_measurement_snapshot", pause_after_snapshot)
    interrupted = reviews.execute_saved_review(store, rid, "first", provider_factory=Provider)
    assert interrupted["status"] == "yielding" and interrupted["api_calls"] == 1
    first_request = Provider.calls[0]
    store.finish_route_measurement_worker(rid, "first")
    resumed = reviews.execute_saved_review(store, rid, "second", provider_factory=Provider)
    assert resumed["status"] == "succeeded", resumed
    assert sum(call == first_request for call in Provider.calls) == 1
    assert resumed["api_calls"] == sum(len(call) - 1 for call in Provider.calls)
    assert resumed["result"]["analysis_result"]["summary"]["in_run_reuse_count"] >= 1


def test_snapshot_access_and_late_write_follow_claim(tmp_path):
    store = SqliteRuntimeStore(tmp_path / "runtime.sqlite")
    row = create(store)
    rid = row["review_id"]
    store.claim_route_measurement_review(rid, "worker")
    assert store.save_review_measurement_snapshot(rid, "worker", "key", {"evidence": 1})
    assert store.get_review_measurement_snapshot(rid, "other", "key") is None
    store.cancel_route_measurement_review(rid)
    assert not store.save_review_measurement_snapshot(rid, "worker", "key", {"evidence": 2})
    assert store.get_review_measurement_snapshot(rid, "worker", "key") is None


@pytest.mark.parametrize("timestamp_field", ["called_at", "measurement_started_at"])
def test_expired_snapshot_is_not_reused(timestamp_field):
    points = source()["prepared_payload"]["original_points"][:2]
    provider = Provider("amap", departure_time=None, api_call_limit=100)
    saved = provider.route(points)
    saved[timestamp_field] = (datetime.now(timezone.utc) - timedelta(seconds=full.CACHE_MAX_AGE_SECONDS + 1)).isoformat()
    wrapped = full.FullReviewProvider(provider, read_snapshot=lambda key: saved, write_snapshot=lambda *_: True, check_active=lambda: None)
    wrapped.route(points)
    assert len(Provider.calls) == 2


@pytest.mark.parametrize("key", ["far_duration_minutes", "stop_service_minutes", "time_window_start", "time_window_end", "service_direction"])
def test_missing_original_parameter_is_not_replaced_by_a_default(key):
    original = source()
    original["result"]["parameters"].pop(key)
    with pytest.raises(ValueError, match="original"):
        full.build_full_review_request(original, requested_by="admin@example.test", request_key="x", provider_call_limit=100)


def test_complete_input_scope_includes_routes_missing_from_partial_old_result():
    original = source()
    original["result"]["routes"] = []
    original["status"] = "failed"
    request = full.build_full_review_request(original, requested_by="admin@example.test", request_key="x", provider_call_limit=100)
    assert [route["route_id"] for route in request["routes"]] == ["R1"]
    assert request["routes"][0]["before"]["drive_duration_s"] is None


@pytest.mark.parametrize("value", [None, "invalid", 1.5, -1, True])
def test_invalid_saved_rider_count_is_not_silently_changed(value):
    original = source()
    original["prepared_payload"]["current_plan"]["stops"][0]["passenger_count"] = value
    with pytest.raises(ValueError, match="passenger count"):
        full.build_full_review_request(original, requested_by="admin@example.test", request_key="x", provider_call_limit=100)


def test_shared_addresses_keep_each_routes_own_students(tmp_path):
    original = source()
    stops = original["prepared_payload"]["current_plan"]["stops"]
    original["prepared_payload"]["current_plan"]["stops"] = stops + [
        {**stop, "route_id": "R2", "passenger_count": 1 if stop["passenger_count"] else 0} for stop in stops]
    store = SqliteRuntimeStore(tmp_path / "runtime.sqlite")
    row = create(store, original)
    completed = reviews.execute_saved_review(store, row["review_id"], "worker", provider_factory=Provider)
    assert completed["status"] == "succeeded", completed
    result = completed["result"]["analysis_result"]
    a = next(stop for stop in result["stops"] if stop["address"] == "A")
    assert a["riders"] == 5
    assert {item["route_id"]: item["passenger_count"] for item in a["occurrences"]} == {"R1": 4, "R2": 1}
    assert {route["route_id"]: route["original_riders"] for route in result["route_window_analysis"]} == {"R1": 10, "R2": 4}


def test_numeric_history_summary_keeps_valid_comparison(tmp_path):
    original = source()
    original["result"]["summary"]["route_only_over_limit_rider_count"] = "10"
    original["prepared_payload"]["current_plan"]["stops"][0]["passenger_count"] = "4.0"
    store = SqliteRuntimeStore(tmp_path / "runtime.sqlite")
    row = create(store, original)
    completed = reviews.execute_saved_review(store, row["review_id"], "worker", provider_factory=Provider)
    assert completed["status"] == "succeeded"
    assert completed["result"]["conclusion_changes"]["route_only_over_limit_rider_count"] == {"before": 10, "after": 3, "delta": -7}
    assert completed["result"]["analysis_result"]["summary"]["direct_over_limit_rider_count"] == 4


def test_full_scope_can_cover_more_than_twenty_routes_but_not_change_inputs(tmp_path):
    original = source()
    stops = original["prepared_payload"]["current_plan"]["stops"]
    original["prepared_payload"]["current_plan"]["stops"] = [
        {**stop, "route_id": f"R{i}"} for i in range(1, 22) for stop in stops]
    store = SqliteRuntimeStore(tmp_path / "runtime.sqlite")
    row = create(store, original)
    assert row["request"]["scope_summary"] == {"route_count": 21, "address_count": 5}
    changed = deepcopy(row["request"])
    changed["full_input"]["analysis_config"]["far_duration_minutes"] = 999
    with pytest.raises(ValueError, match="Source changed"):
        store.create_route_measurement_review(changed, queue_scope="test")


def test_corrected_workbook_reuses_native_statistics_and_adds_differences(tmp_path):
    store = SqliteRuntimeStore(tmp_path / "runtime.sqlite")
    row = create(store)
    completed = reviews.execute_saved_review(store, row["review_id"], "worker", provider_factory=Provider)
    workbook = load_workbook(BytesIO(full.build_full_review_workbook(completed)))
    native = load_workbook(BytesIO(analysis.build_direct_school_workbook({"result": completed["result"]["analysis_result"]})))
    assert "Correction Comparison" in workbook.sheetnames and "Classification Changes" in workbook.sheetnames
    for name in native.sheetnames:
        if name != "Operational Summary":
            assert list(workbook[name].values) == list(native[name].values)
    assert list(workbook["Correction Comparison"].values)[4][1:] == (0, 4, 4)
    assert workbook["Route Outcomes"]["A5"].fill.fgColor == native["Route Outcomes"]["A5"].fill.fgColor
    assert workbook["Route Outcomes"]["A5"].fill.fill_type == "solid"


@pytest.mark.parametrize("change", [
    {"status": "needs_review"}, {"issues": [{"code": "detour"}]}, {"duration_s": 1},
    {"leg_durations_s": [1]}, {"complete": False}, {"provider": "osrm"},
])
def test_incoherent_or_uncertain_measurement_is_not_certified(change):
    points = source()["prepared_payload"]["original_points"][:2]
    provider = Provider("amap", departure_time=None, api_call_limit=100)
    snapshot = provider.route(points)
    assert full._verified(snapshot, 2)
    assert not full._verified({**snapshot, **change}, 2)


def test_full_review_provider_sanitizes_errors(tmp_path):
    store = SqliteRuntimeStore(tmp_path / "runtime.sqlite")
    row = create(store)
    class FailedProvider(Provider):
        def route(self, *args, **kwargs):
            raise RuntimeError("private_key_url_must_not_appear")
    completed = reviews.execute_saved_review(store, row["review_id"], "worker", provider_factory=FailedProvider)
    assert completed["status"] == "needs_review"
    assert "private_key_url" not in str(completed)


def test_source_deletion_cascades_private_measurement_snapshots(tmp_path):
    store = SqliteRuntimeStore(tmp_path / "runtime.sqlite")
    row = create(store)
    rid = row["review_id"]
    store.claim_route_measurement_review(rid, "worker")
    assert store.save_review_measurement_snapshot(rid, "worker", "key", {"evidence": 1})
    assert store.delete_job("source")
    with store.connect() as conn:
        assert conn.execute("SELECT count(*) FROM route_measurement_review_snapshots").fetchone()[0] == 0


def test_full_terminal_result_cannot_omit_routes_or_change_parameters(tmp_path):
    store = SqliteRuntimeStore(tmp_path / "runtime.sqlite")
    row = create(store)
    completed = reviews.execute_saved_review(store, row["review_id"], "worker", provider_factory=Provider)
    changed = deepcopy(completed["result"])
    changed["analysis_result"]["routes"] = []
    with pytest.raises(ValueError, match="every original route"):
        full.validate_full_result(row["request"], changed, terminal=True)
    changed = deepcopy(completed["result"])
    changed["analysis_result"]["parameters"]["far_duration_minutes"] = 999
    with pytest.raises(ValueError, match="original analysis parameters"):
        full.validate_full_result(row["request"], changed, terminal=True)
    changed = deepcopy(completed["result"])
    changed["analysis_result"]["stops"][0]["occurrences"] = []
    with pytest.raises(ValueError, match="each route occurrence"):
        full.validate_full_result(row["request"], changed, terminal=True)


def test_changed_input_digest_stops_before_any_provider_call(tmp_path):
    store = SqliteRuntimeStore(tmp_path / "runtime.sqlite")
    row = create(store)
    row["request"]["full_input"]["analysis_config"]["far_duration_minutes"] = 999
    with pytest.raises(ValueError, match="input or contract changed"):
        full.run_full_review(store, row, "unused", provider_factory=Provider, checkpoint=lambda _: None)
    assert not Provider.calls


def test_zero_dwell_is_preserved_in_corrected_result(tmp_path):
    original = source()
    original["result"]["parameters"]["stop_service_minutes"] = 0
    store = SqliteRuntimeStore(tmp_path / "runtime.sqlite")
    row = create(store, original)
    completed = reviews.execute_saved_review(store, row["review_id"], "worker", provider_factory=Provider)
    assert completed["status"] == "succeeded"
    result = completed["result"]["analysis_result"]
    assert result["parameters"]["stop_service_minutes"] == 0
    assert result["routes"][0]["total_duration_min"] == result["routes"][0]["provider_duration_min"]


def test_schema_seven_upgrade_adds_snapshots_without_changing_review(tmp_path):
    path = tmp_path / "runtime.sqlite"
    store = SqliteRuntimeStore(path)
    row = create(store)
    with store.connect() as conn:
        conn.execute("DROP TABLE route_measurement_review_snapshots")
        conn.execute("UPDATE schema_migrations SET version = 7")
    upgraded = SqliteRuntimeStore(path)
    assert upgraded.get_route_measurement_review(row["review_id"]) == row
    upgraded.claim_route_measurement_review(row["review_id"], "worker")
    assert upgraded.save_review_measurement_snapshot(row["review_id"], "worker", "key", {"value": 1})


def test_full_api_requires_explicit_mode_and_exports_only_completed_result(fixture):
    client, store, schedule = fixture
    store.upsert_job(source())
    body = {"mode": "full_direct_school", "request_key": "full-api", "provider_call_limit": 100, "confirm_provider_calls": True}
    base = "/api/jobs/source/measurement-reviews"
    assert client.post(base, json={**body, "route_keys": ["current_plan:R1"]}).status_code == 422
    row = client.post(base, json=body).json()
    assert row["request"]["mode"] == "full_direct_school", row
    url = f"{base}/{row['review_id']}/export"
    assert client.get(url).status_code == 409
    reviews.execute_saved_review(store, row["review_id"], "worker", provider_factory=Provider)
    assert client.get(url, headers={"X-BRP-User-Email": "stranger@example.test"}).status_code == 403
    response = client.get(url, headers={"X-BRP-User-Email": "owner@example.test"})
    assert response.status_code == 200 and response.headers["content-disposition"].startswith("attachment")
    assert "Correction Comparison" in load_workbook(BytesIO(response.content)).sheetnames
