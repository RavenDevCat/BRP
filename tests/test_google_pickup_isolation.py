from copy import deepcopy
from io import BytesIO

from openpyxl import load_workbook
import pytest

from test_google_final_validation import client, response, body_points
from test_google_timing_entrypoints import context, CONFIG
from test_direct_school_analysis import prepared_payload, fake_osrm
import direct_school_analysis as direct
import google_final_validation as google


def two_routes():
    payload = prepared_payload()
    stops = payload['current_plan']['stops']
    stops.extend([{**deepcopy(stops[1]), 'route_id': 'R2', 'stop_sequence': 1},
                  {**deepcopy(stops[2]), 'route_id': 'R2', 'stop_sequence': 2}])
    return payload


@pytest.mark.parametrize('failure_scope', ['direct', 'current', 'both'])
def test_bad_pickup_does_not_poison_other_routes(context, monkeypatch, failure_scope):
    payload = two_routes()
    monkeypatch.setattr(direct, '_osrm_leg', fake_osrm)
    original = context.route
    def route(points, **kwargs):
        has_far = any(p['address'] == 'Far stop' for p in points)
        current = kwargs.get('dwell_s') is not None
        if has_far and (failure_scope == 'both' or current == (failure_scope == 'current')):
            raise google.ValidationUnavailable('google_pickup_snap_mismatch', details={
                'address': 'Far stop', 'point_index': 0, 'endpoint': 'start', 'snap_distance_m': 152.77})
        return original(points, **kwargs)
    monkeypatch.setattr(context, 'route', route)
    result = direct.run_direct_school_analysis(payload, CONFIG, timing_context=context)
    assert result['status'] == 'partial' and result['measurement_attempts_complete'] is True
    routes = {r['route_id']: r for r in result['routes']}
    assert routes['R2']['status'] == 'resolved' and routes['R2']['total_duration_min'] > 0
    rows = {r['address']: r for r in result['stops']}
    good = next(c for c in rows['Near stop']['route_contexts'] if c['route_id'] == 'R2')
    assert good['operational_category'] == 'within_limit'
    assert good['estimated_current_ride_min'] > 0
    assert result['operational_conclusion']['final']['all_measured_routes_within_window'] is False
    assert result['operational_conclusion']['data_review']['rider_count'] > 0
    if failure_scope != 'direct':
        assert routes['R1'].get('total_duration_min') is None
        assert routes['R1'].get('geometry') in (None, [])
    if failure_scope != 'direct':
        expected = sum(s['passenger_count'] for s in payload['current_plan']['stops'] if s['route_id'] == 'R1')
        assert routes['R1']['riders'] == expected
        assert routes['R1']['stop_count'] == sum(not s.get('is_depot', False) for s in payload['current_plan']['stops'] if s['route_id'] == 'R1')
        window = next(r for r in result['route_window_analysis'] if r['route_id'] == 'R1')
        assert window['original_riders'] == expected
        exported = next(r for r in direct._route_outcome_rows(result) if r[0] == 'R1')
        assert exported[1] == expected
        assert exported[3] is None

    if failure_scope != 'current':
        assert rows['Far stop'].get('direct_duration_min') is None
    record = {'job_id': 'partial-test', 'status': 'succeeded', 'result': result}
    book = load_workbook(BytesIO(direct.build_direct_school_workbook(record, include_diagnostics=True)))
    assert book['Operational Summary']['F3'].value == 'partial'
    assert book['Operational Summary']['A4'].value == 'Partial forecast / 部分预测'
    assert book['Data Quality'].max_row > 4
    for status in ['failed', 'running', 'canceled']:
        with pytest.raises(ValueError, match='not available'):
            direct.build_direct_school_workbook({**record, 'status': status})
    result.pop('measurement_attempts_complete')
    with pytest.raises(ValueError, match='not available'):
        direct.build_direct_school_workbook(record)


@pytest.mark.parametrize('error', ['google_http_403', 'google_transport_failed', 'google_budget_cap', 'canceled'])
def test_global_errors_still_abort_and_keep_checkpoint(context, monkeypatch, error):
    monkeypatch.setattr(direct, '_osrm_leg', fake_osrm)
    def fail(*args, **kwargs):
        raise google.ValidationUnavailable(error)
    monkeypatch.setattr(context, 'route', fail)
    checkpoints = []
    with pytest.raises(google.ValidationUnavailable, match=error):
        direct.run_direct_school_analysis(two_routes(), CONFIG, timing_context=context, checkpoint=checkpoints.append)
    assert checkpoints and checkpoints[-1].get('measurement_attempts_complete') is not True


def test_removal_recheck_snap_failure_does_not_approve_removal(context, monkeypatch):
    monkeypatch.setattr(direct, '_osrm_leg', fake_osrm)
    def fail(*args, **kwargs):
        raise google.ValidationUnavailable('google_pickup_snap_mismatch', details={'address': 'Far stop'})
    monkeypatch.setattr(direct, '_measure_active_route', fail)
    config = {**CONFIG, 'time_window_start': '07:40', 'time_window_end': '08:00'}
    result = direct.run_direct_school_analysis(two_routes(), config, timing_context=context)
    windows = {r['route_id']: r for r in result['route_window_analysis']}
    assert windows['R1']['status'] == 'data_review'
    assert windows['R2']['status'] == 'within_window'
    assert result['operational_conclusion']['additional_removal']['rider_count'] == 0
    assert result['operational_conclusion']['final']['all_measured_routes_within_window'] is False
