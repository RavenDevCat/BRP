from copy import deepcopy
from datetime import datetime, timedelta, timezone
from types import SimpleNamespace
import pytest
import google_final_validation as google
from google_timing_fallback import HybridValidationSession, LocalTimingFallback, provenance

NOW=datetime(2026,9,28,tzinfo=timezone.utc)
POINTS=[(31.2,121.4+i*0.01) for i in range(4)]

def leg(a,b,duration=60):
    return {'duration_s':duration,'distance_m':900,'start':a,'end':b,
            'geometry':[list(reversed(a)),list(reversed(b))],'coordinate_system':'WGS84'}

class Native:
    def __init__(self):
        self.state={'api_calls':0}
    def route(self,points):
        self.state['api_calls']+=1
        a,b=[(p['plot_lat'],p['plot_lng']) for p in points]
        return {'status':'verified','complete':True,'provider':'amap','called_at':NOW.isoformat(),
                'legs':[leg(a,b,120)]}

@pytest.fixture
def make(monkeypatch):
    import google_geocoding
    monkeypatch.setattr(google_geocoding,'GoogleGeocodeResolver',lambda *a,**kw:SimpleNamespace())
    def build(error='google_pickup_snap_mismatch'):
        requests=[]
        def route(points,when):
            requests.append((deepcopy(points),when))
            if len(points)>2 or tuple(points[0])==POINTS[1]:
                raise google.ValidationUnavailable(error,details={'leg_index':1 if len(points)>2 else 0,'point_index':1})
            return [leg(*points)]
        client=SimpleNamespace(budget_id='unit',check_canceled=lambda:None,route=route,
            can_afford=lambda *a:True,protected_calls=0,request_headroom=0)
        native=Native()
        fallback=LocalTimingFallback(provider=native,now=lambda:NOW)
        session=HybridValidationSession(client,fallback=fallback)
        session.bind_pickups([{'plot_lat':a,'plot_lng':b} for a,b in POINTS])
        return session,requests,native
    return build

def test_only_failed_edge_falls_back_and_suffix_uses_updated_clock(make):
    session,requests,native=make()
    result=session.measure(POINTS,NOW,[30,30,30,0])
    assert result.drive_s==240 and result.dwell_s==90
    assert [x.get('timing_provider') for x in result.legs]==['google_routes','amap','google_routes']
    assert requests[-1][1]==NOW+timedelta(seconds=270)
    assert native.state['api_calls']==1
    assert provenance(result.legs)['forecast_complete'] is False
    assert provenance(result.legs)['fallback_leg_indexes']==[1]
    assert 'not a complete Google' in provenance(result.legs)['timing_note']

@pytest.mark.parametrize('error',['google_http_403','google_budget_cap','google_transport_failed',
    'google_duration_reconciliation','google_response_invalid','google_pickup_identity_unresolved'])
def test_systemic_and_identity_failures_never_trigger_native_fallback(make,error):
    session,requests,native=make(error)
    with pytest.raises(google.ValidationUnavailable,match=error):
        session.measure(POINTS,NOW,[30,30,30,0])
    assert native.state['api_calls']==0

def test_failed_batch_recovers_edges_without_retrying_known_bad_edge(make):
    session,requests,native=make()
    result=session.measure(POINTS,NOW,[0]*4)
    assert len(requests)==3
    assert len(result.legs)==3 and result.drive_s==240
    assert native.state['api_calls']==1

def test_current_traffic_reuse_does_not_claim_future_forecast(make):
    session,requests,native=make()
    session.measure(POINTS,NOW,[30,30,30,0])
    result=session.measure(POINTS,NOW+timedelta(hours=1),[30,30,30,0])
    assert native.state['api_calls']==1
    assert result.legs[1]['time_basis']=='current_traffic'
    assert result.legs[1]['provider_called_at']==NOW.isoformat()
    assert len(requests)==5

def test_native_endpoint_or_metrics_failure_does_not_become_success(make):
    session,requests,native=make()
    original=native.route
    def wrong(points):
        result=original(points)
        result['legs'][0]['geometry'][0][0]+=0.01
        return result
    native.route=wrong
    with pytest.raises(google.ValidationUnavailable,match='google_local_timing_endpoint_mismatch'):
        session.measure(POINTS,NOW,[30,30,30,0])

def test_mixed_join_remeasures_only_incident_edges(make):
    session,requests,native=make()
    original=session.client.route
    def shifted(points,when):
        result=original(points,when)
        if tuple(points[0])==POINTS[0]:
            result[0]['end']=(POINTS[1][0]+0.0005,POINTS[1][1])
        return result
    session.client.route=shifted
    result=session.measure(POINTS,NOW,[30,30,30,0])
    assert [x.get('timing_provider') for x in result.legs]==['amap','amap','google_routes']
    assert native.state['api_calls']==2
    assert requests[-1][1]==NOW+timedelta(seconds=330)

def test_native_cache_expires_without_sliding_refresh(make):
    session,requests,native=make()
    session.measure(POINTS,NOW,[30,30,30,0])
    session.fallback.now=lambda:NOW+timedelta(minutes=11)
    session.measure(POINTS,NOW+timedelta(minutes=11),[30,30,30,0])
    assert native.state['api_calls']==2

def test_identity_binding_is_required(make):
    session,requests,native=make()
    session.fallback.points={}
    with pytest.raises(google.ValidationUnavailable,match='google_fallback_identity_unavailable'):
        session.measure(POINTS,NOW,[30,30,30,0])
    assert native.state['api_calls']==0


def test_arrival_inversion_rechecks_google_with_mixed_duration(make):
    session,requests,native=make()
    latest=NOW+timedelta(hours=1)
    result=session.validate(POINTS,[30,30,30,0],NOW,latest,120,True)
    assert result.arrival==latest
    assert result.departure==latest-timedelta(seconds=330)
    assert native.state['api_calls']==1
    assert len(requests)==5
    assert provenance(result.legs)['forecast_complete'] is False


def test_boarding_after_fallback_keeps_google_only_provenance(make):
    session,requests,native=make()
    result=session.measure(POINTS,NOW,[30,30,30,0])
    assert provenance(result.legs[2:])['forecast_complete'] is True
    assert provenance(result.legs[:2])['forecast_complete'] is False


def test_public_export_keeps_fallback_disclosure_without_diagnostics():
    from direct_school_analysis import _address_measurement_rows
    mixed=provenance([{'timing_provider':'amap'}])
    rows=_address_measurement_rows({'stops':[{'address':'sample', 'route_evidence':mixed,
        'route_contexts':[{'route_id':'R1',**mixed}], 'operational_category':'within_limit'}]})
    assert 'AMap current traffic' in rows[0][-2]
    assert 'not complete Google future prediction' in rows[0][-1]


def test_actual_public_workbook_discloses_mixed_timing():
    from io import BytesIO
    from openpyxl import load_workbook
    from direct_school_analysis import build_direct_school_workbook
    mixed={**provenance([{'timing_provider':'amap'}]),'status':'verified','complete':True,
           'duration_s':600,'distance_m':2000}
    row={'stop_key':'sample','address':'Sample','riders':1,'route_ids':['R1'],
         'operational_category':'within_limit','direct_duration_min':10,'direct_distance_km':2,
         'estimated_current_ride_min':10,'provider_status':'resolved','route_evidence':mixed,
         'route_contexts':[{'route_id':'R1',**mixed,'estimated_current_ride_min':10}]}
    record={'job_id':'mixed-export','status':'succeeded','result':{
        'analysis_version':6,'provider':'google_routes','status':'complete','service_direction':'To School',
        'parameters':{'far_duration_minutes':60},'stops':[row],'routes':[],'route_window_analysis':[]}}
    workbook=load_workbook(BytesIO(build_direct_school_workbook(record)))
    assert 'Route Evidence' not in workbook.sheetnames
    sheet=workbook['Address Measurements']
    assert 'AMap current traffic' in sheet['K5'].value
    assert 'not complete Google future prediction' in sheet['L5'].value
    assert any('not complete Google future prediction' in str(cell.value)
               for row in workbook['Operational Summary'] for cell in row)
