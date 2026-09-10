"""Distance diagnostics cannot veto complete measurements or hide real failures."""
from copy import deepcopy
from io import BytesIO
import importlib
import json
from pathlib import Path
import sys

import pytest
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path[:0] = [str(ROOT / "apps"), str(ROOT / "apps/backend")]
import route_evidence as evidence
import planner_core as core
import direct_school_analysis as analysis
from direct_school_review import present_direct_school_result
from route_measurement_view import route_display_metrics, measurement_note
from test_route_evidence import A, B, C, PRECISE_PICKUP, leg


class WarningPlanner:
    AMAP_KEY = "test"
    AMAP_ROUTING_LIMITER = object()
    MAX_ROUTE_DURATION_SECONDS = 7200
    OSRM_BASE_URL = "test"

    def amap_request_json(self, endpoint, params, limiter):
        points = [tuple(reversed(tuple(map(float, p.split(","))))) for p in
                  [params["origin"], *filter(None, params.get("waypoints", "").split(";")), params["destination"]]]
        durations = [1200, 2400] if len(points) == 3 else [600] * (len(points) - 1)
        steps = [{"step_distance": "1400", "cost": {"duration": str(seconds)},
                  "polyline": ";".join(f"{p[1]},{p[0]}" for p in points[i:i + 2]),
                  "navi": {"assistant_action": "\u5230\u8fbe\u9014\u7ecf\u5730" if i < len(durations) - 1 else "\u5230\u8fbe\u76ee\u7684\u5730"}}
                 for i, seconds in enumerate(durations)]
        return {"route": {"paths": [{"distance": str(1400 * len(durations)),
                                    "cost": {"duration": str(sum(durations))}, "steps": steps}]}}

    @staticmethod
    def point_osrm_lat_lng(point, **kwargs):
        return point["lat"], point["lng"]

    @staticmethod
    def osrm_request_json(*args):
        return {"routes": [{"duration": 200, "distance": 800, "geometry": {"coordinates": []},
                            "legs": [{"duration": 100, "distance": 400}] * 2}]}


def points():
    return [{"lat": p[0], "lng": p[1], "provider": "amap", "country": "China", **PRECISE_PICKUP,
             "passenger_count": 0 if i == 0 else 1, "is_depot": i == 0}
            for i, p in enumerate([C, A, B])]


class GapPlanner(WarningPlanner):
    def amap_request_json(self, endpoint, params, limiter):
        payload = super().amap_request_json(endpoint, params, limiter)
        path = payload["route"]["paths"][0]
        step = path["steps"][0]
        a, b = [list(map(float, pair.split(","))) for pair in step["polyline"].split(";")]
        before = [x + (y - x) * .1 for x, y in zip(a, b)]
        after = [x + (y - x) * .9 for x, y in zip(a, b)]
        first = deepcopy(step)
        first["polyline"] = ";".join(",".join(map(str, p)) for p in [a, before])
        first["navi"] = {}
        step["polyline"] = ";".join(",".join(map(str, p)) for p in [after, b])
        first["cost"]["duration"] = step["cost"]["duration"] = str(float(step["cost"]["duration"]) / 2)
        first["step_distance"] = step["step_distance"] = "700"
        path["steps"].insert(0, first)
        return payload


def test_native_warning_metrics_geometry_and_calls_are_preserved():
    state = {"api_call_limit": 1, "expected_leg_distances_m": [400, 400]}
    result = evidence.measure_amap_route(WarningPlanner(), [A, B, C], {}, state)
    assert result["status"] == "verified" and not result["issues"]
    assert {w["code"] for w in result["warnings"]} == evidence.DISTANCE_WARNING_CODES
    assert result["duration_s"] == 3600 and result["distance_m"] == 2800
    assert result["leg_durations_s"] == [1200, 2400]
    assert state["api_calls"] == 1
    assert result["geometry_segments"] == [leg["geometry"] for leg in result["legs"]]


@pytest.mark.parametrize("end,expected", [("08:00", "passed"), ("07:30", "failed")])
@pytest.mark.parametrize("scenario_name", ["Strict Plan", "Protected Plan"])
def test_final_gate_uses_time_not_warning_and_preserves_assignment(monkeypatch, end, expected, scenario_name):
    monkeypatch.setattr(core, "infer_traffic_location", lambda _: ("CHINA", "Shanghai"))
    monkeypatch.setattr(core, "FINAL_ROUTE_TRAFFIC_VERIFICATION_ENABLED", True)
    route = {"route_id": "R1", "nodes": [1, 2, 0], "time_s": 200, "distance_m": 800,
             "stop_service_time_s": 120, "vehicle_id": 4,
             "leg_details": [{"duration_s": 100, "distance_m": 400}] * 2}
    scenario = {"routes": [route]}
    config = core.PlannerConfig(service_direction="To School", stop_service_minutes=1,
                                time_window_start="06:30", time_window_end=end)
    gate = core.attach_final_route_traffic_gate(WarningPlanner(), scenario, points(), config, [], scenario_name)
    assert gate["status"] == expected and scenario["traffic_feasible"] == (expected == "passed")
    assert route["nodes"] == [1, 2, 0] and route["vehicle_id"] == 4
    assert route["final_route_traffic_gate"]["verified_total_duration_s"] == 3720
    rows, unavailable = core._scenario_time_impact_rows(scenario, points(), config)
    assert unavailable == 0 and [r["offset_s"] for r in rows] == [-3720, -2460]
    assert route_display_metrics(route)["duration_s"] == 3720
    assert "Suspected detour" not in measurement_note(route)


@pytest.mark.parametrize("problem", ["snap", "missing", "shorter", "continuity"])
def test_distance_warning_never_clears_integrity_failure(problem):
    def fetch(planner, pair):
        result = {**leg(planner, pair), "distance_m": 1400}
        if problem == "snap":
            result["geometry"][0][0] += .01
        elif problem == "missing" and pair[0] == B:
            result["geometry"] = []
        elif problem == "shorter" and pair[0] == B:
            result["distance_m"] = 1
        elif problem == "continuity" and pair[0] == B:
            result["geometry"][0][0] += .001
        return result
    result = evidence._measure_adjacent_route(object(), [A, B, C], {},
        {"api_call_limit": 10, "expected_leg_distances_m": [400, 400]}, fetch_leg=fetch)
    assert result["status"] in {"needs_review", "unavailable"}
    assert result["issues"] and result["warnings"]
    assert route_display_metrics({"route_evidence": result, "stop_service_time_s": 120})["duration_s"] is None


def test_fleet_and_insert_keep_warning_separate_from_feasibility(monkeypatch):
    service = importlib.import_module("backend_service")
    api = importlib.import_module("api_app")
    monkeypatch.setattr(core, "load_legacy_planner", lambda: WarningPlanner())
    monkeypatch.setattr(service, "load_legacy_planner", lambda: WarningPlanner())
    ordered = points()[1:] + points()[:1]
    routing = service._client_module("demand_routing")
    route = {"cluster_id": "G01", "order": [1, 2, 0], "ordered_points": ordered,
             "selected_vehicle": {"capacity": 42}, "leg_details": [{"distance_m": 400}] * 2}
    plan = {"school": {"country": "China"}, "summary": {"service_direction": "to_school", "max_route_duration_minutes": 90},
            "routes": [route], "route_rows": [{"cluster_id": "G01"}]}
    service._attach_fleet_route_measurements(plan, routing)
    assert route["final_route_traffic_gate"]["status"] == "passed"
    assert "Suspected detour" not in measurement_note(route)
    measured = api._insert_route_measurement(ordered, "China", {})
    assert measured["provider_verified"] and measured["duration_s"] == 3600
    assert not measured["display_geometry_message"]
    assert "amap_final_validation_unavailable" not in measured["warnings"]


@pytest.mark.parametrize("scheduled", [False, True])
@pytest.mark.parametrize("gap", [False, True])
def test_runner_classification_removal_and_export_use_warning_measurements(monkeypatch, scheduled, gap, tmp_path):
    from test_direct_school_analysis import prepared_payload
    runner = importlib.import_module("backend_job_runner")
    prepared = prepared_payload()
    for point, pos in zip(prepared["original_points"], [C, A, B]):
        point.update(lat=pos[0], lng=pos[1], plot_lat=pos[0], plot_lng=pos[1], provider="amap", **PRECISE_PICKUP)
    monkeypatch.setattr(core, "load_legacy_planner", lambda: GapPlanner() if gap else WarningPlanner())
    monkeypatch.setattr(analysis, "_osrm_leg", lambda *args: {"distance_m": 400, "duration_s": 100})
    state = {"record": {"job_id": "warning-runner", "status": "queued", "prepared_payload": prepared,
        "scheduled_start_at": "2026-09-10T23:00:00+00:00" if scheduled else None,
        "metadata": {"job_kind": "direct_school_analysis", "scheduled_job": scheduled,
                     "analysis_config": {"far_duration_minutes": 45}}}}
    monkeypatch.setattr(runner, "_load_job", lambda _: deepcopy(state["record"]))
    monkeypatch.setattr(runner, "_save_job", lambda record: state.update(record=deepcopy(record)))
    monkeypatch.setattr(runner, "_release_concurrency_slot", lambda: None)
    monkeypatch.setattr(sys, "argv", ["backend_job_runner.py", "warning-runner"])
    assert runner.main() == 0
    result = state["record"]["result"]
    assert result["analysis_version"] == 8 and result["routes"][0]["status"] == "resolved"
    assert result["summary"]["route_measurement_verified_count"] == 1
    assert result["summary"]["route_measurement_review_count"] == 0
    assert result["measurement_warnings"]
    if gap:
        assert result["routes"][0]["route_evidence"]["geometry_diagnostics"]
        assert len(result["routes"][0]["route_evidence"]["geometry_segments"]) == 3
        assert len(result["stops"][0]["direct_geometry_segments"]) == 2
    rows = {row["address"]: row for row in result["stops"]}
    assert rows["Far stop"]["estimated_current_ride_min"] == 62
    assert rows["Near stop"]["estimated_current_ride_min"] == 41
    assert rows["Far stop"]["operational_category"] == "route_only_over_limit"
    assert rows["Near stop"]["operational_category"] == "within_limit"
    assert result["route_window_analysis"][0]["status"] != "data_review"
    book = load_workbook(BytesIO(analysis.build_direct_school_workbook(state["record"])))
    assert "Route Warnings" not in book.sheetnames
    assert "Route Evidence" not in book.sheetnames
    diagnostic = load_workbook(BytesIO(analysis.build_direct_school_workbook(state["record"], include_diagnostics=True)))
    assert diagnostic["Route Warnings"].max_row > 4
    before = deepcopy(result)
    assert present_direct_school_result(result) == result and result == before
    (tmp_path / "warning-record.json").write_text(json.dumps({**state["record"], "multi_day": {"run_count": 0, "stops": []}}))
    (tmp_path / "warning-export.xlsx").write_bytes(analysis.build_direct_school_workbook(state["record"]))


def test_old_review_result_is_not_silently_promoted():
    raw = {"status": "partial", "routes": [{"route_id": "R2", "status": "needs_review", "route_evidence": {
        "provider": "amap", "status": "needs_review", "complete": True,
        "source": "amap_continuous_waypoint_legs", "point_count": 3,
        "leg_durations_s": [1200, 2400], "leg_distances_m": [1400, 1400], "duration_s": 3600,
        "distance_m": 2800, "issues": [{"leg_index": 0, "code": "provider_distance_disagreement"}]}}],
        "stops": [], "operational_conclusion": {"final": {"all_measured_routes_within_window": False}}}
    before = deepcopy(raw)
    shown = present_direct_school_result(raw)
    assert shown["routes"][0]["status"] == "needs_review"
    assert shown["operational_conclusion"] == raw["operational_conclusion"] and raw == before
    assert shown["measurement_warnings"] == []


def test_automatic_budget_does_not_skip_complete_warning_route(monkeypatch):
    service = importlib.import_module("backend_service")
    monkeypatch.setattr(service, "save_json_object", lambda *args, **kwargs: None)
    details = {"minutes": 10}
    service._attach_current_plan_amap_budget_details(WarningPlanner(), details,
        {"stops": [{"country": "China", "city": "Shanghai"}]}, points(), {"R1": [1, 2, 0]})
    assert details["amap_route_measured_count"] == 1
    assert details["minutes"] == 62
    assert details["amap_route_duration_minutes"] == 62


@pytest.mark.parametrize("module_name", ["BusingProblem", "client_runtime"])
def test_legacy_map_export_does_not_reconnect_intra_leg_gap(module_name, tmp_path):
    service = importlib.import_module("backend_service")
    module = core.load_legacy_planner() if module_name == "BusingProblem" else service._client_module("client_runtime")
    saved = evidence.measure_amap_route(GapPlanner(), [A, B, C], {}, {"api_call_limit": 1})
    route = {"route_id": "R1", "nodes": [0, 1, 2], "time_s": 3720, "distance_m": 2800,
             "stop_service_time_s": 120, "vehicle_id": 1, "bus_type_name": "Bus", "load": 2,
             "bus_capacity": 42, "route_evidence": saved}
    original = deepcopy(route)
    stops = [{"plot_lat": p[0], "plot_lng": p[1], "address": f"Stop {i}", "passenger_count": 1}
             for i, p in enumerate([A, B, C])]
    file = tmp_path / f"{module_name}.html"
    module.render_map(stops, [route], str(file))
    markup = file.read_text()
    assert markup.count("L.polyline(") == 6
    assert "native_step_geometry_gap" not in markup and "Suspected detour" not in markup
    assert route == original
