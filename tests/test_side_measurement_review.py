from concurrent.futures import ThreadPoolExecutor
from copy import deepcopy
from io import BytesIO
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

import side_measurement_review as side
import measurement_reviews as reviews
from runtime_store_sqlite import SqliteRuntimeStore
from test_fleet_measurement_views import source as fleet_plan
from test_full_measurement_review import Provider
from test_route_insert_measurements import source as insert_map, action, measured
from test_measurement_review_api import fixture
import api_app as api


def fleet_source(direction="to_school"):
    return {"run_id": "same-id", "owner_email": "owner@example.test", "global_plan_result": fleet_plan(direction)}


def insert_source(monkeypatch, direction="to_school"):
    monkeypatch.setattr(api, "_insert_route_measurement", lambda points, *_: measured(points))
    plan, data = api._insert_build_selected_plan(insert_map(direction=direction), [action()], country="China", constraints={},
        suggested_config={"service_direction": direction, "stop_service_minutes": 0,
                          "time_window_start": "06:00", "time_window_end": "07:00"})
    return {"run_id": "same-id", "owner_email": "owner@example.test", "route_insert_result": {"summary": {},
        "scenarios": [{"id": "recommended", "selected_plan": plan, "selected_map_data": data}]}}


def request(source, tool, key="review"):
    return side.build_side_review_request(source, tool_key=tool, requested_by="admin@example.test",
        request_key=key, provider_call_limit=100)


def create(store, source, tool, key="review"):
    store.upsert_side_tool_run(tool, source)
    return store.create_route_measurement_review(request(source, tool, key), queue_scope="test")


@pytest.mark.parametrize("direction", ["to_school", "from_school"])
@pytest.mark.parametrize("tool", ["fleet_planner", "route_insert_advisor"])
def test_native_review_preserves_source_order_counts_and_saved_dwell(tmp_path, monkeypatch, tool, direction):
    source = fleet_source(direction) if tool == "fleet_planner" else insert_source(monkeypatch, direction)
    store = SqliteRuntimeStore(tmp_path / "runtime.sqlite")
    row = create(store, source, tool)
    before = store.get_side_tool_run(tool, source["run_id"])
    result = reviews.execute_saved_review(store, row["review_id"], "worker", provider_factory=Provider)
    assert result["status"] == "succeeded", result["result"]
    assert result["source_job_id"] is None and result["source_tool_key"] == tool
    output = result["result"]
    assert output["scope"] == side.MODES[tool] + "_result"
    assert output["time_window_revalidated"] and not output["student_classification_recomputed"]
    assert store.get_side_tool_run(tool, source["run_id"]) == before
    assert not store.list_route_measurement_reviews("same-id")
    assert len(store.list_route_measurement_reviews(None, source_tool_key=tool, source_run_id="same-id")) == 1
    side.validate_side_result(row["request"], output, terminal=True)
    for scope in row["request"]["routes"]:
        zero = next(p for p in scope["points"] if "waypoint" in p["address"].lower())
        assert zero["passenger_count"] == 0 and zero["is_depot"] is False
    native = side.corrected_native_record(result)
    assert native["run_id"] == "same-id" and native["measurement_review_id"] == row["review_id"]
    if tool == "fleet_planner":
        plan = native["global_plan_result"]
        assert plan["map_data"]["routes"][0]["duration_s"] == 2120
        assert plan["routes"][0]["stop_service_time_s"] == 120
        assert plan["routes"][0]["final_route_traffic_gate"]["status"] == "failed"
        assert plan["workbook_base64"]
    else:
        plan = native["route_insert_result"]["selected_plan"]
        assert plan["affected_routes"][0]["measurement_inputs"]["base_stop_service_time_s"] == 0
        assert plan["affected_routes"][0]["selected_duration_s"] == 3000


def test_native_parent_identity_and_concurrent_idempotency(tmp_path, monkeypatch):
    store = SqliteRuntimeStore(tmp_path / "runtime.sqlite")
    first, other = fleet_source(), insert_source(monkeypatch)
    store.upsert_side_tool_run("fleet_planner", first)
    store.upsert_side_tool_run("route_insert_advisor", other)
    payload = request(first, "fleet_planner")
    with ThreadPoolExecutor(max_workers=4) as pool:
        records = list(pool.map(lambda _: store.create_route_measurement_review(deepcopy(payload), queue_scope="test"), range(4)))
    assert len({row["review_id"] for row in records}) == 1
    second = create(store, other, "route_insert_advisor")
    assert second["review_id"] != records[0]["review_id"]
    changed = deepcopy(payload)
    changed["source_run_id"] = "missing"
    with pytest.raises(ValueError):
        store.create_route_measurement_review(changed)
    changed = deepcopy(payload)
    changed["routes"][0]["points"][0]["lat"] += .1
    with pytest.raises(ValueError):
        store.create_route_measurement_review(changed)
    changed = deepcopy(first)
    changed["global_plan_result"]["routes"][0]["ordered_points"][0]["student_count"] = 8
    store.upsert_side_tool_run("fleet_planner", changed)
    with pytest.raises(ValueError):
        store.create_route_measurement_review(payload)


@pytest.mark.parametrize("change", ["dwell", "coordinates", "window", "school", "capacity"])
def test_missing_fleet_inputs_disable_full_review_without_dropping_risk(change):
    original = fleet_source()
    route = original["global_plan_result"]["routes"][0]
    if change == "dwell":
        route.pop("stop_service_time_s")
    elif change == "coordinates":
        route["ordered_points"][0].pop("coordinate_system")
    elif change == "window":
        original["global_plan_result"]["summary"].pop("max_route_duration_minutes")
    elif change == "school":
        original["global_plan_result"]["school"] = {"address": "Different school"}
    else:
        route["selected_vehicle"].pop("student_capacity")
    before = deepcopy(original)
    risk = side.side_risk_summary(original, "fleet_planner", "admin@example.test")
    assert not risk["full_review"]["available"] and risk["routes"]
    assert original == before


def test_legacy_insert_missing_inputs_remains_explicitly_unavailable(monkeypatch):
    original = insert_source(monkeypatch)
    original["route_insert_result"]["scenarios"][0]["selected_plan"]["affected_routes"][0].pop("measurement_inputs")
    risk = side.side_risk_summary(original, "route_insert_advisor", "admin@example.test")
    assert not risk["full_review"]["available"] and risk["routes"]


@pytest.mark.parametrize("direction", ["to_school", "from_school"])
@pytest.mark.parametrize("shape", ["native", "legacy", "duplicate"])
def test_fleet_recovers_only_coordinate_provenance_from_saved_geocodes(direction, shape):
    source = fleet_source(direction)
    points = source["global_plan_result"]["routes"][0]["ordered_points"]
    school_index = len(points) - 1 if direction == "to_school" else 0
    geocodes = deepcopy(points)
    source["geocode_result"] = ({"points": geocodes} if shape == "legacy" else {
        "school": geocodes[school_index],
        "demand_points": [point for index, point in enumerate(geocodes) if index != school_index],
    })
    if shape == "duplicate":
        source["geocode_result"]["points"] = deepcopy(geocodes)
    for point in points:
        point.pop("coordinate_system")
    for point in geocodes:
        point["student_count"] = 999
    before = deepcopy(source)
    scopes = side.source_scopes(source, "fleet_planner")
    assert not scopes[0]["input_issues"]
    assert [point["coordinate_system"] for point in scopes[0]["points"]] == [point["coordinate_system"] for point in geocodes]
    assert [point["passenger_count"] for point in scopes[0]["points"]] == [point["student_count"] for point in points]
    assert source == before


@pytest.mark.parametrize("change", ["address", "lat", "conflicting_system", "missing_system"])
def test_fleet_does_not_guess_missing_coordinate_provenance(change):
    source = fleet_source()
    point = source["global_plan_result"]["routes"][0]["ordered_points"][0]
    saved = deepcopy(point)
    source["geocode_result"] = {"demand_points": [saved]}
    point.pop("coordinate_system")
    if change == "address":
        saved["address"] += " other gate"
    elif change == "lat":
        saved["lat"] += .001
    elif change == "missing_system":
        saved.pop("coordinate_system")
    else:
        source["geocode_result"]["points"] = [{**saved, "coordinate_system": "GCJ02" if saved["coordinate_system"] == "WGS84" else "WGS84"}]
    before = deepcopy(source)
    assert "source_stop_coordinates_missing" in side.source_scopes(source, "fleet_planner")[0]["input_issues"]
    assert source == before


def test_review_results_cannot_switch_native_parent_or_change_student_count(tmp_path):
    store = SqliteRuntimeStore(tmp_path / "runtime.sqlite")
    row = create(store, fleet_source(), "fleet_planner")
    store.claim_route_measurement_review(row["review_id"], "worker")
    output = side.run_side_review(store, row, "worker", provider_factory=Provider, checkpoint=lambda _: None)
    wrong = deepcopy(output)
    wrong["source_tool_key"] = "route_insert_advisor"
    with pytest.raises(ValueError):
        store.save_route_measurement_review(row["review_id"], "worker", wrong, terminal=True)
    wrong = deepcopy(output)
    wrong["native_result"]["global_plan_result"]["routes"][0]["ordered_points"][0]["student_count"] = 99
    with pytest.raises(ValueError):
        store.save_route_measurement_review(row["review_id"], "worker", wrong, terminal=True)


def test_native_budget_pause_snapshot_and_cascade(tmp_path):
    store = SqliteRuntimeStore(tmp_path / "runtime.sqlite")
    source = fleet_source()
    source["route_preview_result"] = deepcopy(source["global_plan_result"])
    row = create(store, source, "fleet_planner")
    store.claim_route_measurement_review(row["review_id"], "first")
    def checkpoint(_):
        assert store.pause_route_measurement_review(row["review_id"])
        raise reviews.ReviewClaimLost("paused")
    with pytest.raises(reviews.ReviewClaimLost):
        side.run_side_review(store, row, "first", provider_factory=Provider, checkpoint=checkpoint)
    used = store.get_route_measurement_review(row["review_id"])["api_calls"]
    assert used == 2
    with store.connect() as conn:
        assert conn.execute("SELECT count(*) FROM route_measurement_review_snapshots").fetchone()[0] == 1
    store.finish_route_measurement_worker(row["review_id"], "first")
    assert store.resume_route_measurement_review(row["review_id"])
    result = reviews.execute_saved_review(store, row["review_id"], "second", provider_factory=Provider)
    assert result["status"] == "succeeded", result
    assert result["api_calls"] == used
    assert store.delete_side_tool_run("fleet_planner", "same-id")
    assert store.get_route_measurement_review(row["review_id"]) is None
    with store.connect() as conn:
        assert conn.execute("SELECT count(*) FROM route_measurement_review_snapshots").fetchone()[0] == 0


def test_native_api_access_scope_and_confirmation(fixture, monkeypatch):
    client, store, schedule = fixture
    store.upsert_side_tool_run("fleet_planner", fleet_source())
    monkeypatch.setattr(api.backend_service, "FLEET_PLANNER_HISTORY_STORE", SimpleNamespace(get=lambda key: store.get_side_tool_run("fleet_planner", key)))
    monkeypatch.setattr(api.backend_service, "ROUTE_INSERT_ADVISOR_HISTORY_STORE", SimpleNamespace(get=lambda key: store.get_side_tool_run("route_insert_advisor", key)))
    base = "/api/fleet-planner/history/same-id/measurement-reviews"
    body = {"mode": "full_fleet", "request_key": "api", "provider_call_limit": 10, "confirm_provider_calls": True}
    for change in ({"confirm_provider_calls": False}, {"provider_call_limit": "10"}, {"route_keys": ["x"]}, {"coordinates": []}):
        assert client.post(base, json={**body, **change}).status_code == 422
    assert client.post(base, json=body, headers={"X-BRP-User-Email": "owner@example.test"}).status_code == 403
    created = client.post(base, json=body)
    assert created.status_code == 200, created.text
    row = created.json()
    assert client.post(base, json=body).json() == row
    assert not {"worker_token", "job_slot_path", "queue_scope", "points"}.intersection(row)
    detail = base + "/" + row["review_id"]
    for path in (base, detail, base.replace("measurement-reviews", "measurement-risk")):
        assert client.get(path, headers={"X-BRP-User-Email": "stranger@example.test"}).status_code == 403
        assert client.get(path, headers={"X-BRP-User-Email": "owner@example.test"}).status_code == 200
    store.upsert_side_tool_run("fleet_planner", {**fleet_source(), "run_id": "other"})
    assert client.get(detail.replace("same-id", "other")).status_code == 404
    assert client.get(detail + "/native-result").status_code == 409
    assert client.post(detail + "/actions/cancel").json()["status"] == "canceled"
    assert not store.list_route_measurement_reviews("same-id")


def test_v8_migration_preserves_active_lease_budget_and_snapshot(tmp_path, monkeypatch):
    path = tmp_path / "legacy.sqlite"
    with monkeypatch.context() as patch:
        patch.setattr(SqliteRuntimeStore, "_migrate_review_sources", lambda *_: None)
        old = SqliteRuntimeStore(path)
        old.upsert_job({"job_id": "job", "status": "succeeded"})
        with old.connect() as conn:
            conn.execute("INSERT INTO route_measurement_reviews(review_id,source_job_id,request_key,status,created_at,worker_token,worker_pid,job_slot_path,api_calls,request_json) VALUES('review','job','key','running','now','lease',123,'slot',7,'{}')")
            conn.execute("INSERT INTO route_measurement_review_snapshots VALUES('review','snapshot','{\"saved\":true}')")
    new = SqliteRuntimeStore(path)
    row = new.get_route_measurement_review("review")
    assert row["api_calls"] == 7 and row["worker_pid"] == 123 and row["job_slot_path"] == "slot"
    assert row["source_job_id"] == "job" and row["source_tool_key"] is None
    assert new.get_review_measurement_snapshot("review", "lease", "snapshot") == {"saved": True}
    with new.connect() as conn:
        assert not conn.execute("PRAGMA foreign_key_check").fetchall()
        conn.execute("DELETE FROM jobs WHERE job_id='job'")
        assert conn.execute("SELECT count(*) FROM route_measurement_review_snapshots").fetchone()[0] == 0


@pytest.mark.parametrize("language", ["en", "zh", "ko"])
def test_native_report_is_read_only_localized_and_preserves_unknowns(tmp_path, language):
    store = SqliteRuntimeStore(tmp_path / "runtime.sqlite")
    source = fleet_source()
    store.upsert_side_tool_run("fleet_planner", source)
    payload = request(source, "fleet_planner")
    payload["provider_call_limit"] = 1
    row = store.create_route_measurement_review(payload)
    result = reviews.execute_saved_review(store, row["review_id"], "worker", provider_factory=Provider)
    assert result["status"] == "needs_review", result
    before = deepcopy(result)
    workbook = load_workbook(BytesIO(side.build_side_review_workbook(result, language)))
    sheet = workbook.worksheets[1]
    assert sheet["E2"].value is None and sheet["H2"].value is None
    assert sheet["I2"].value == 3
    assert sheet.freeze_panes == "A2"
    assert workbook.worksheets[2]["H2"].value is None
    assert result == before


def test_migration_unknown_dependent_table_rolls_back(tmp_path, monkeypatch):
    path = tmp_path / "legacy.sqlite"
    with monkeypatch.context() as patch:
        patch.setattr(SqliteRuntimeStore, "_migrate_review_sources", lambda *_: None)
        old = SqliteRuntimeStore(path)
        old.initialize()
        with old.connect() as conn:
            conn.execute("CREATE TABLE future_dependency (review_id TEXT REFERENCES route_measurement_reviews(review_id))")
    new = SqliteRuntimeStore(path)
    with pytest.raises(RuntimeError, match="Unknown review dependent"):
        new.initialize()
    with new.connect() as conn:
        assert "source_tool_key" not in {row["name"] for row in conn.execute("PRAGMA table_info(route_measurement_reviews)")}
        assert not conn.execute("SELECT name FROM sqlite_master WHERE name='route_measurement_reviews_v9'").fetchall()


def test_native_api_final_result_and_export_follow_the_authorized_source(fixture, monkeypatch):
    client, store, schedule = fixture
    store.upsert_side_tool_run("fleet_planner", fleet_source())
    store.upsert_side_tool_run("route_insert_advisor", insert_source(monkeypatch))
    monkeypatch.setattr(api.backend_service, "FLEET_PLANNER_HISTORY_STORE", SimpleNamespace(get=lambda key: store.get_side_tool_run("fleet_planner", key)))
    monkeypatch.setattr(api.backend_service, "ROUTE_INSERT_ADVISOR_HISTORY_STORE", SimpleNamespace(get=lambda key: store.get_side_tool_run("route_insert_advisor", key)))
    base = "/api/fleet-planner/history/same-id/measurement-reviews"
    body = {"mode": "full_fleet", "request_key": "final", "provider_call_limit": 10, "confirm_provider_calls": True}
    row = client.post(base, json=body).json()
    result = reviews.execute_saved_review(store, row["review_id"], "worker", provider_factory=Provider)
    assert result["status"] == "succeeded"
    monkeypatch.setattr(api.backend_service, "FreshRouteProvider", lambda *_a, **_k: pytest.fail("No provider during reads/exports"))
    path = base + "/" + row["review_id"]
    assert client.get(path.replace("fleet-planner", "route-insert-advisor")).status_code == 404
    native = client.get(path + "/native-result").json()
    assert native["measurement_review_id"] == row["review_id"] and native["run_id"] == "same-id"
    assert native["global_plan_result"]["map_data"]["routes"][0]["duration_s"] == 2120
    for language in ("en", "zh", "ko"):
        response = client.get(path + "/export", params={"language": language})
        assert response.status_code == 200
        workbook = load_workbook(BytesIO(response.content))
        assert workbook.worksheets[1]["E2"].value == pytest.approx(35.33)
    assert client.get(path + "/export", headers={"X-BRP-User-Email": "stranger@example.test"}).status_code == 403
